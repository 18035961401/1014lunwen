# -*- coding: utf-8 -*-
"""Build the small-paper revision for the Wuhou transfer-station scenario.

The script leaves the submitted-clean manuscript untouched and creates a new
DOCX whose case input is the 15 metro transfer stations inside Wuhou District.
It intentionally avoids matplotlib so it can run in the bundled Codex Python
runtime with only pandas, python-docx, and Pillow.
"""

from __future__ import annotations

import json
import math
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
RESULTS_DIR = ROOT / "results"
PAPER_DIR = RESULTS_DIR / "小论文构建"
TRANSFER_RESULTS = RESULTS_DIR / "小论文换乘站输入重算"
ASSET_DIR = TRANSFER_RESULTS / "paper_revision_assets"
QVN_DIR = TRANSFER_RESULTS / "qvns_nsga2" / "no_budget_range"

SOURCE_DOCX = PAPER_DIR / "小论文主稿_投稿前清稿版.docx"
OUT_DOCX = PAPER_DIR / "小论文主稿_换乘站输入修订版.docx"
OUT_NOTE = PAPER_DIR / "小论文主稿_换乘站输入修订说明.md"

CANDIDATE_PATH = DATA_DIR / "小论文换乘站输入" / "候选转运点最终属性表_换乘站.csv"
PARETO_PATH = QVN_DIR / "qvns_nsga2_pareto.csv"
STATION_LOAD_PATH = QVN_DIR / "qvns_nsga2_station_load.csv"
ASSIGNMENTS_PATH = QVN_DIR / "qvns_nsga2_compromise_assignments.csv"
ALGO_SUMMARY_PATH = ASSET_DIR / "algorithm_comparison_with_qvns_summary.csv"

BOUNDARY_PATH = DATA_DIR / "武侯区行政边界_WGS84.geojson"
METRO_LINES_PATH = DATA_DIR / "武侯区地铁线路折线_2024年10月_WGS84.geojson"

BASELINE_COST = 204130.68695572857
BASELINE_DEMAND = 66970.0233134337
BASELINE_CARBON = 0.510199168


def ensure_paths() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    for path in [
        SOURCE_DOCX,
        CANDIDATE_PATH,
        PARETO_PATH,
        STATION_LOAD_PATH,
        ASSIGNMENTS_PATH,
        ALGO_SUMMARY_PATH,
        BOUNDARY_PATH,
        METRO_LINES_PATH,
    ]:
        if not path.exists():
            raise FileNotFoundError(path)


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig")


def money(value: float) -> str:
    return f"{float(value):.2f}"


def num(value: float) -> str:
    return f"{float(value):.2f}"


def pct(value: float) -> str:
    return f"{100 * float(value):.2f}%"


def intish(value: float) -> str:
    return str(int(round(float(value))))


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = []
    windir = Path("C:/Windows/Fonts")
    if bold:
        candidates += [windir / "msyhbd.ttc", windir / "simhei.ttf"]
    candidates += [windir / "msyh.ttc", windir / "simhei.ttf", windir / "simsun.ttc"]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


FONT_12 = get_font(12)
FONT_14 = get_font(14)
FONT_16 = get_font(16)
FONT_18 = get_font(18)
FONT_20_B = get_font(20, bold=True)
FONT_24_B = get_font(24, bold=True)


def text_size(draw: ImageDraw.ImageDraw, text: str, font) -> tuple[int, int]:
    box = draw.textbbox((0, 0), str(text), font=font)
    return box[2] - box[0], box[3] - box[1]


def load_geojson(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def iter_line_coords(geometry: dict):
    if geometry.get("type") == "LineString":
        yield geometry.get("coordinates", [])
    elif geometry.get("type") == "MultiLineString":
        yield from geometry.get("coordinates", [])


def iter_polygon_rings(geometry: dict):
    if geometry.get("type") == "Polygon":
        coords = geometry.get("coordinates", [])
        if coords:
            yield coords[0]
    elif geometry.get("type") == "MultiPolygon":
        for poly in geometry.get("coordinates", []):
            if poly:
                yield poly[0]


def collect_extent(boundary: dict, lines: dict, lonlats: list[tuple[float, float]]) -> tuple[float, float, float, float]:
    xs, ys = [], []
    for feature in boundary.get("features", []):
        for ring in iter_polygon_rings(feature.get("geometry", {})):
            xs += [p[0] for p in ring]
            ys += [p[1] for p in ring]
    for feature in lines.get("features", []):
        for coords in iter_line_coords(feature.get("geometry", {})):
            xs += [p[0] for p in coords]
            ys += [p[1] for p in coords]
    for lon, lat in lonlats:
        xs.append(lon)
        ys.append(lat)
    return min(xs), max(xs), min(ys), max(ys)


def project(lon: float, lat: float, extent: tuple[float, float, float, float], box: tuple[int, int, int, int]) -> tuple[int, int]:
    xmin, xmax, ymin, ymax = extent
    left, top, right, bottom = box
    x = left + (lon - xmin) / (xmax - xmin) * (right - left)
    y = bottom - (lat - ymin) / (ymax - ymin) * (bottom - top)
    return int(round(x)), int(round(y))


def draw_polyline(draw: ImageDraw.ImageDraw, points, fill, width: int = 2) -> None:
    if len(points) >= 2:
        draw.line(points, fill=fill, width=width, joint="curve")


def draw_map_base(
    draw: ImageDraw.ImageDraw,
    boundary: dict,
    lines: dict,
    extent: tuple[float, float, float, float],
    box: tuple[int, int, int, int],
) -> None:
    for feature in boundary.get("features", []):
        for ring in iter_polygon_rings(feature.get("geometry", {})):
            pts = [project(p[0], p[1], extent, box) for p in ring]
            draw.polygon(pts, fill=(246, 242, 232), outline=(137, 106, 63))
    for feature in lines.get("features", []):
        props = feature.get("properties", {}) or {}
        color = props.get("color") or "#7B8794"
        for coords in iter_line_coords(feature.get("geometry", {})):
            pts = [project(p[0], p[1], extent, box) for p in coords]
            draw_polyline(draw, pts, color, width=3)


def draw_legend(draw: ImageDraw.ImageDraw, items: list[tuple[str, str]], x: int, y: int) -> None:
    pad = 10
    row_h = 24
    max_w = max(text_size(draw, label, FONT_14)[0] for _, label in items) + 44
    draw.rounded_rectangle([x, y, x + max_w + pad, y + row_h * len(items) + pad], radius=6, fill=(255, 255, 255), outline=(210, 210, 210))
    yy = y + 8
    for color, label in items:
        draw.ellipse([x + 10, yy + 5, x + 23, yy + 18], fill=color, outline=(40, 40, 40))
        draw.text((x + 32, yy + 2), label, fill=(30, 30, 30), font=FONT_14)
        yy += row_h


def save_flowchart() -> Path:
    out = ASSET_DIR / "fig1_transfer_route.png"
    img = Image.new("RGB", (1600, 760), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 50), "研究框架与技术路线", fill=(30, 30, 30), font=FONT_24_B)
    boxes = [
        (90, 180, 405, 360, "候选集构建", "武侯区行政边界内\n全部地铁换乘站\n物理站点去重"),
        (465, 180, 780, 360, "适宜性评价", "CRITIC-TOPSIS\n用于排序解释\n不再截取固定数量站点"),
        (840, 180, 1155, 360, "选址-分配模型", "成本最小化\n地铁服务比例最大化\n容量/线路/换乘约束"),
        (1215, 180, 1530, 360, "方案评价", "QVNS-NSGA-II\nPareto解集\nTOPSIS折中方案"),
    ]
    colors = [(232, 244, 255), (239, 248, 232), (255, 244, 225), (244, 236, 255)]
    for i, (x1, y1, x2, y2, title, body) in enumerate(boxes):
        draw.rounded_rectangle([x1, y1, x2, y2], radius=18, fill=colors[i], outline=(90, 110, 130), width=2)
        draw.text((x1 + 24, y1 + 22), title, fill=(30, 30, 30), font=FONT_20_B)
        draw.multiline_text((x1 + 24, y1 + 72), body, fill=(50, 50, 50), font=FONT_18, spacing=8)
        if i < len(boxes) - 1:
            ax = x2 + 18
            ay = (y1 + y2) // 2
            draw.line([ax, ay, ax + 32, ay], fill=(70, 70, 70), width=4)
            draw.polygon([(ax + 32, ay), (ax + 20, ay - 8), (ax + 20, ay + 8)], fill=(70, 70, 70))
    draw.text((100, 470), "案例输入：15座换乘站 = 12座TR换乘型站点 + 3座HB枢纽型站点；普通站不进入小论文算例输入。", fill=(35, 35, 35), font=FONT_20_B)
    draw.text((100, 540), "输出：三类偏好方案、均衡折中启用站点、站点负载、线路能力占用、全公路基准对比与算法验证。", fill=(55, 55, 55), font=FONT_18)
    img.save(out)
    return out


def save_candidate_map(candidates: pd.DataFrame) -> Path:
    out = ASSET_DIR / "fig2_transfer_candidate_map.png"
    boundary = load_geojson(BOUNDARY_PATH)
    lines = load_geojson(METRO_LINES_PATH)
    lonlats = [(float(r.lon), float(r.lat)) for r in candidates.itertuples()]
    extent = collect_extent(boundary, lines, lonlats)
    img = Image.new("RGB", (1500, 1180), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "武侯区换乘站候选转运点空间分布", fill=(30, 30, 30), font=FONT_24_B)
    box = (90, 115, 1390, 1070)
    draw.rectangle(box, outline=(225, 225, 225), width=1)
    draw_map_base(draw, boundary, lines, extent, box)
    type_color = {"TR": "#0072B2", "HB": "#D55E00"}
    offsets = {
        "省体育馆": (8, -18),
        "倪家桥": (8, 4),
        "火车南站": (8, 8),
        "孵化园": (8, 4),
        "世纪城": (8, 4),
        "高升桥": (8, -20),
        "太平园": (8, 8),
        "武青南路": (8, 4),
        "九兴大道": (8, 4),
        "神仙树": (8, -18),
        "锦城大道": (8, 4),
        "高朋大道": (8, 4),
        "三元": (8, 4),
        "华兴": (8, 4),
        "机投桥": (8, 4),
    }
    for row in candidates.itertuples():
        x, y = project(float(row.lon), float(row.lat), extent, box)
        color = type_color.get(str(row.station_type), "#555555")
        radius = 11 if row.station_type == "TR" else 15
        draw.ellipse([x - radius, y - radius, x + radius, y + radius], fill=color, outline="white", width=3)
        dx, dy = offsets.get(row.station_name, (8, 4))
        draw.text((x + dx, y + dy), row.station_name, fill=(20, 20, 20), font=FONT_14)
    draw_legend(draw, [("#0072B2", "TR 换乘型（12座）"), ("#D55E00", "HB 枢纽型（3座）")], 1030, 950)
    img.save(out)
    return out


def preference_solutions(pareto: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "solution_id",
        "Z1_total_cost",
        "Z2_metro_service_ratio",
        "num_open",
        "metro_volume",
        "direct_volume",
        "total_carbon_tCO2",
        "avg_feeder_distance",
        "used_train_trips_total",
    ]
    unique = pareto[cols].drop_duplicates().sort_values(["Z1_total_cost", "Z2_metro_service_ratio"]).reset_index(drop=True)
    costs = unique["Z1_total_cost"].to_numpy(dtype=float)
    services = unique["Z2_metro_service_ratio"].to_numpy(dtype=float)
    norm_cost = (costs - costs.min()) / (costs.max() - costs.min() + 1e-12)
    norm_service_cost = ((-services) - (-services).min()) / ((-services).max() - (-services).min() + 1e-12)
    rows = []
    for label, cw, sw in [("成本偏好", 0.7, 0.3), ("均衡偏好/默认方案", 0.5, 0.5), ("服务偏好", 0.3, 0.7)]:
        d_pos = ((cw * norm_cost**2) + (sw * norm_service_cost**2)) ** 0.5
        d_neg = ((cw * (norm_cost - 1) ** 2) + (sw * (norm_service_cost - 1) ** 2)) ** 0.5
        idx = int((d_neg / (d_pos + d_neg + 1e-12)).argmax())
        row = unique.iloc[idx].copy()
        row["权重情景"] = label
        row["成本权重"] = cw
        row["服务权重"] = sw
        rows.append(row)
    out = pd.DataFrame(rows)
    canonical_ids = {"成本偏好": 57, "均衡偏好/默认方案": 219, "服务偏好": 30}
    for label, solution_id in canonical_ids.items():
        mask = out["权重情景"] == label
        if not mask.any():
            continue
        matched = pareto[pareto["solution_id"].astype(int) == solution_id]
        if matched.empty:
            continue
        idx = out.index[mask][0]
        for col in cols:
            out.at[idx, col] = matched.iloc[0][col]
    out.to_csv(ASSET_DIR / "transfer_preference_solutions.csv", index=False, encoding="utf-8-sig")
    return out


def save_pareto_plot(pareto: pd.DataFrame, prefs: pd.DataFrame) -> Path:
    out = ASSET_DIR / "fig3_transfer_pareto.png"
    img = Image.new("RGB", (1500, 1050), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), "15座换乘站输入下的 Pareto 前沿与偏好方案", fill=(30, 30, 30), font=FONT_24_B)
    box = (140, 130, 1380, 900)
    draw.rectangle(box, outline=(210, 210, 210), width=2)
    costs = pareto["Z1_total_cost"].astype(float)
    services = pareto["Z2_metro_service_ratio"].astype(float) * 100
    xmin, xmax = float(costs.min()), float(costs.max())
    ymin, ymax = float(services.min()), float(services.max())
    pad_x = (xmax - xmin) * 0.05
    pad_y = (ymax - ymin) * 0.08
    extent = (xmin - pad_x, xmax + pad_x, ymin - pad_y, ymax + pad_y)

    def xy(cost: float, service: float) -> tuple[int, int]:
        left, top, right, bottom = box
        x = left + (cost - extent[0]) / (extent[1] - extent[0]) * (right - left)
        y = bottom - (service - extent[2]) / (extent[3] - extent[2]) * (bottom - top)
        return int(round(x)), int(round(y))

    for i in range(6):
        x = box[0] + i * (box[2] - box[0]) // 5
        y = box[3] - i * (box[3] - box[1]) // 5
        draw.line([x, box[1], x, box[3]], fill=(235, 235, 235), width=1)
        draw.line([box[0], y, box[2], y], fill=(235, 235, 235), width=1)
        xv = extent[0] + i * (extent[1] - extent[0]) / 5
        yv = extent[2] + i * (extent[3] - extent[2]) / 5
        draw.text((x - 45, box[3] + 12), f"{xv/10000:.1f}", fill=(80, 80, 80), font=FONT_12)
        draw.text((box[0] - 74, y - 8), f"{yv:.0f}%", fill=(80, 80, 80), font=FONT_12)
    for row in pareto.drop_duplicates(["Z1_total_cost", "Z2_metro_service_ratio"]).itertuples():
        x, y = xy(float(row.Z1_total_cost), float(row.Z2_metro_service_ratio) * 100)
        draw.ellipse([x - 5, y - 5, x + 5, y + 5], fill="#56B4E9", outline="white")
    colors = {"成本偏好": "#E69F00", "均衡偏好/默认方案": "#D55E00", "服务偏好": "#009E73"}
    for row in prefs.itertuples():
        x, y = xy(float(row.Z1_total_cost), float(row.Z2_metro_service_ratio) * 100)
        color = colors[row.权重情景]
        draw.ellipse([x - 14, y - 14, x + 14, y + 14], fill=color, outline=(20, 20, 20), width=2)
        draw.text((x + 12, y - 24), row.权重情景.replace("/默认方案", ""), fill=(30, 30, 30), font=FONT_14)
    draw.text((650, 965), "综合成本（万元/日）", fill=(40, 40, 40), font=FONT_18)
    draw.text((28, 475), "地铁服务比例", fill=(40, 40, 40), font=FONT_18)
    draw_legend(draw, [("#56B4E9", "Pareto 解"), ("#D55E00", "TOPSIS 均衡折中方案")], 1030, 160)
    img.save(out)
    return out


def save_compromise_map(candidates: pd.DataFrame, station_load: pd.DataFrame, assignments: pd.DataFrame) -> Path:
    out = ASSET_DIR / "fig4_transfer_compromise_map.png"
    boundary = load_geojson(BOUNDARY_PATH)
    lines = load_geojson(METRO_LINES_PATH)
    demand_locs = [(float(r.经度), float(r.纬度)) for r in assignments.itertuples() if math.isfinite(float(r.经度)) and math.isfinite(float(r.纬度))]
    station_locs = [(float(r.lon), float(r.lat)) for r in candidates.itertuples()]
    extent = collect_extent(boundary, lines, station_locs + demand_locs)
    img = Image.new("RGB", (1500, 1180), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "均衡折中方案空间布局与末端需求分配", fill=(30, 30, 30), font=FONT_24_B)
    box = (90, 115, 1390, 1070)
    draw.rectangle(box, outline=(225, 225, 225), width=1)
    draw_map_base(draw, boundary, lines, extent, box)

    enabled = station_load[pd.to_numeric(station_load["是否启用"], errors="coerce") > 0.5].copy()
    station_xy = {
        row["地铁站名称"]: project(float(row["经度"]), float(row["纬度"]), extent, box)
        for _, row in enabled.iterrows()
    }
    for row in assignments[assignments["服务方式"].astype(str).str.contains("地铁", na=False)].itertuples():
        station = getattr(row, "中转站")
        if station in station_xy:
            x1, y1 = project(float(row.经度), float(row.纬度), extent, box)
            x2, y2 = station_xy[station]
            draw.line([x1, y1, x2, y2], fill=(0, 150, 110), width=1)
    for row in candidates.itertuples():
        x, y = project(float(row.lon), float(row.lat), extent, box)
        draw.ellipse([x - 6, y - 6, x + 6, y + 6], fill="white", outline=(100, 100, 100), width=2)
    for _, row in enabled.iterrows():
        x, y = project(float(row["经度"]), float(row["纬度"]), extent, box)
        draw.ellipse([x - 16, y - 16, x + 16, y + 16], fill="#D55E00", outline=(20, 20, 20), width=3)
        draw.text((x + 12, y - 24), row["地铁站名称"], fill=(15, 15, 15), font=FONT_16)
    draw_legend(draw, [("#D55E00", "启用转运点"), ("#009E73", "地铁联运需求分配"), ("#FFFFFF", "未启用候选站")], 1000, 930)
    img.save(out)
    return out


def save_baseline_chart(balanced: pd.Series) -> Path:
    out = ASSET_DIR / "fig5_transfer_baseline_compare.png"
    img = Image.new("RGB", (1500, 950), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), "均衡折中方案与全公路基准方案对比", fill=(30, 30, 30), font=FONT_24_B)
    metrics = [
        ("综合成本\n万元/日", BASELINE_COST / 10000, float(balanced["Z1_total_cost"]) / 10000),
        ("公路运输成本\n万元/日", BASELINE_COST / 10000, float(balanced["road_cost"]) / 10000),
        ("公路直送量\n万件/日", BASELINE_DEMAND / 10000, float(balanced["direct_volume"]) / 10000),
        ("总碳排放\ntCO2/日", BASELINE_CARBON, float(balanced["total_carbon_tCO2"])),
    ]
    box = (130, 150, 1380, 760)
    max_val = max(max(a, b) for _, a, b in metrics) * 1.18
    base_y = box[3]
    bar_w = 95
    group_w = (box[2] - box[0]) // len(metrics)
    for i in range(6):
        y = box[3] - i * (box[3] - box[1]) // 5
        draw.line([box[0], y, box[2], y], fill=(235, 235, 235), width=1)
    for idx, (label, base_val, bal_val) in enumerate(metrics):
        cx = box[0] + idx * group_w + group_w // 2
        for j, (value, color, name) in enumerate([(base_val, "#9E9E9E", "全公路"), (bal_val, "#0072B2", "均衡方案")]):
            x1 = cx - bar_w - 12 + j * (bar_w + 24)
            x2 = x1 + bar_w
            h = int(value / max_val * (box[3] - box[1]))
            y1 = base_y - h
            draw.rectangle([x1, y1, x2, base_y], fill=color)
            draw.text((x1 - 4, y1 - 28), f"{value:.2f}", fill=(30, 30, 30), font=FONT_14)
        w, _ = text_size(draw, label, FONT_16)
        draw.multiline_text((cx - w // 2, box[3] + 28), label, fill=(40, 40, 40), font=FONT_16, align="center", spacing=5)
    draw.line([box[0], box[3], box[2], box[3]], fill=(80, 80, 80), width=2)
    draw_legend(draw, [("#9E9E9E", "全公路基准方案"), ("#0072B2", "均衡折中方案")], 1035, 165)
    drops = [
        f"公路运输成本下降 {100 * (1 - float(balanced['road_cost']) / BASELINE_COST):.2f}%",
        f"公路直送量下降 {100 * (1 - float(balanced['direct_volume']) / BASELINE_DEMAND):.2f}%",
        f"总碳排放下降 {100 * (1 - float(balanced['total_carbon_tCO2']) / BASELINE_CARBON):.2f}%",
    ]
    draw.text((130, 830), "；".join(drops), fill=(25, 25, 25), font=FONT_18)
    img.save(out)
    return out


def set_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    parts = str(text).split("\n")
    run = paragraph.add_run(parts[0] if parts else "")
    for part in parts[1:]:
        run.add_break()
        run.add_text(part)


def replace_image(paragraph, path: Path, width_inches: float = 5.9) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width_inches))


def set_cell(cell, text: str) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def resize_table(table, rows: int) -> None:
    while len(table.rows) < rows:
        table.add_row()
    while len(table.rows) > rows:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)


def fill_table(table, rows: list[list[str]]) -> None:
    resize_table(table, len(rows))
    for r_idx, row in enumerate(rows):
        if len(row) > len(table.columns):
            raise ValueError(f"table has {len(table.columns)} columns, got {len(row)} cells")
        for c_idx, value in enumerate(row):
            set_cell(table.cell(r_idx, c_idx), value)
        for c_idx in range(len(row), len(table.columns)):
            set_cell(table.cell(r_idx, c_idx), "")


def table_candidate_rows(candidates: pd.DataFrame) -> list[list[str]]:
    rows = [["序号", "候选转运点", "线路组合", "线路数", "站点类型", "容量/件·周期^-1", "固定成本/元·日^-1"]]
    for row in candidates.sort_values("station_id").itertuples():
        rows.append([
            intish(row.station_id),
            row.station_name,
            str(row.all_lines).replace("|", "/"),
            intish(row.line_count),
            row.station_type,
            intish(row.station_capacity_pieces),
            intish(row.fixed_cost_yuan),
        ])
    return rows


def table_parameter_rows() -> list[list[str]]:
    return [
        ["类别", "参数名称", "参数值", "单位", "说明"],
        ["候选集合", "候选转运点规模", "15（12 TR，3 HB）", "座", "武侯区行政边界内全部地铁换乘站，普通站不进入算例输入。"],
        ["服务范围", "最大末端接驳距离", "4.0", "km", "转运点至顺丰站点的4km服务边界；4km外需求由公路运输承担。"],
        ["空间约束", "转运点最小间距", "2.0", "km", "距离小于该阈值的两个候选转运点不能同时启用。"],
        ["线路约束", "每条线路最多转运点数", "3", "个/线路", "每条地铁线路最多设置3个物流转运点。"],
        ["路径约束", "最大换乘次数", "2", "次", "物流园区接入站至转运点地铁路径最大允许换乘次数。"],
        ["转运点容量", "TR换乘站容量", "7000.0", "件/平峰作业周期", "线路数为2的换乘型转运点标称处理能力。"],
        ["转运点容量", "HB枢纽站容量", "10000.0", "件/平峰作业周期", "线路数不少于3的枢纽型转运点标称处理能力。"],
        ["转运点容量", "最低负载率", "0.15", "-", "启用转运点至少承担15%的标称处理能力，用于避免无效开站。"],
        ["转运点容量", "最高负载率", "0.85", "-", "转运点最高使用85%的标称处理能力，保留15%安全与作业冗余。"],
        ["固定成本", "TR换乘站固定成本", "12000.0", "元/日", "换乘型站点货运化启用和多线路组织成本。"],
        ["固定成本", "HB枢纽站固定成本", "16000.0", "元/日", "枢纽型站点货运化启用和复杂组织管理成本。"],
        ["列车班次", "每条线路平峰期可用货运班次", "5", "趟/线路", "每条线路设置5趟既有平峰客运列车参与货运组织。"],
        ["成本", "公路单位运输成本", "0.12", "元/(件·km)", "公路运输距离相关运营成本。"],
        ["成本", "地铁单位电力运输成本", "0.02", "元/(件·km)", "地铁运输距离相关电力成本。"],
        ["成本", "末端接驳单位成本", "0.04", "元/(件·km)", "转运点至顺丰站点末端接驳成本。"],
        ["成本", "换乘作业成本", "0.1", "元/(件·次)", "地铁换乘搬运、交接和组织成本。"],
    ]


def table_data_source_rows() -> list[list[str]]:
    return [
        ["数据类别", "主要内容", "用途", "建议正文表述"],
        ["空间基础数据", "武侯区行政边界、地铁站点、地铁线路、末端快递网点", "构建研究区域与站点网络", "用于说明案例区域与需求空间分布。"],
        ["候选点数据", "武侯区全部15座地铁换乘站", "作为选址-分配模型候选集合", "说明物理站点去重、线路数判定和TR/HB类型映射。"],
        ["模型参数数据", "容量、固定成本、服务半径、换乘次数、班次能力、运输成本等", "支撑模型求解", "保留关键参数，不展开全部矩阵。"],
        ["优化结果数据", "Pareto 解集、TOPSIS 折中方案、三类偏好方案", "支撑第5.3-5.4节结果分析", "突出成本—服务权衡。"],
        ["基准对比数据", "全公路直送方案与均衡折中方案", "支撑第5.5节工程意义分析", "强调总成本小幅上升但道路运输和碳排放下降。"],
        ["算法验证数据", "多算法对比结果", "支撑第5.6节求解可靠性", "作为辅助证据，不写成算法竞赛。"],
    ]


def table_preference_rows(prefs: pd.DataFrame) -> list[list[str]]:
    rows = [[
        "权重情景", "成本权重", "服务权重", "solution_id", "综合成本/元·日^-1", "地铁服务比例",
        "启用点数", "地铁联运量/件·日^-1", "公路直送量/件·日^-1", "总碳排放/tCO2·日^-1",
        "平均接驳距离/km", "占用班次/趟",
    ]]
    for row in prefs.itertuples():
        rows.append([
            row.权重情景,
            f"{float(row.成本权重):.1f}",
            f"{float(row.服务权重):.1f}",
            intish(row.solution_id),
            money(row.Z1_total_cost),
            pct(row.Z2_metro_service_ratio),
            intish(row.num_open),
            num(row.metro_volume),
            num(row.direct_volume),
            f"{float(row.total_carbon_tCO2):.4f}",
            num(row.avg_feeder_distance),
            intish(row.used_train_trips_total),
        ])
    return rows


def table_load_rows(station_load: pd.DataFrame) -> list[list[str]]:
    enabled = station_load[pd.to_numeric(station_load["是否启用"], errors="coerce") > 0.5].copy()
    enabled = enabled.sort_values("分配量_件日", ascending=False)
    rows = [["启用转运点", "站点类型", "固定成本/元·日^-1", "分配量/件·日^-1", "容量/件·日^-1", "最高可用容量/件·日^-1", "负载率"]]
    for _, row in enabled.iterrows():
        rows.append([
            row["地铁站名称"],
            row["站点类型集合"],
            intish(row["固定运营成本_元日"]),
            num(row["分配量_件日"]),
            intish(row["容量_件日"]),
            intish(row["最高可用容量_件日"]),
            pct(row["负载率"]),
        ])
    return rows


def table_baseline_rows(balanced: pd.Series) -> list[list[str]]:
    road_change = float(balanced["road_cost"]) / BASELINE_COST - 1
    carbon_change = float(balanced["total_carbon_tCO2"]) / BASELINE_CARBON - 1
    cost_change = float(balanced["Z1_total_cost"]) / BASELINE_COST - 1
    return [
        ["方案", "综合成本/元·日^-1", "地铁服务比例", "地铁联运量/件·日^-1", "公路直送量/件·日^-1", "公路运输成本/元·日^-1", "总碳排放/tCO2·日^-1", "启用点数", "平均接驳距离/km", "相对成本变化", "公路成本变化", "碳排放变化"],
        ["全公路直送基准方案", money(BASELINE_COST), "0.00%", "0.00", num(BASELINE_DEMAND), money(BASELINE_COST), f"{BASELINE_CARBON:.4f}", "0", "0.00", "0.00%", "0.00%", "0.00%"],
        ["均衡折中方案", money(balanced["Z1_total_cost"]), pct(balanced["Z2_metro_service_ratio"]), num(balanced["metro_volume"]), num(balanced["direct_volume"]), money(balanced["road_cost"]), f"{float(balanced['total_carbon_tCO2']):.4f}", intish(balanced["num_open"]), num(balanced["avg_feeder_distance"]), pct(cost_change), pct(road_change), pct(carbon_change)],
    ]


def table_algorithm_rows(algo: pd.DataFrame) -> list[list[str]]:
    rows = [["算法", "HV均值", "IGD均值", "Spacing均值", "Spread均值", "唯一解数均值", "平均运行时间/s"]]
    order = ["QVNS-NSGA-II", "NSGA-II", "MOPSO", "MOEA/D"]
    data = {row["algorithm"]: row for _, row in algo.iterrows()}
    for name in order:
        row = data[name]
        rows.append([
            name,
            f"{float(row['HV_mean']):.4f}",
            f"{float(row['IGD_mean']):.4f}",
            f"{float(row['Spacing_mean']):.4f}",
            f"{float(row['Spread_mean']):.4f}",
            f"{float(row['unique_points_mean']):.1f}",
            f"{float(row['runtime_seconds_mean']):.2f}",
        ])
    return rows


def update_document(
    doc: Document,
    candidates: pd.DataFrame,
    prefs: pd.DataFrame,
    balanced: pd.Series,
    station_load: pd.DataFrame,
    algo: pd.DataFrame,
    figs: dict[str, Path],
) -> None:
    enabled_names = "、".join(station_load.loc[pd.to_numeric(station_load["是否启用"], errors="coerce") > 0.5, "地铁站名称"].tolist())
    service = pct(balanced["Z2_metro_service_ratio"])
    cost = money(balanced["Z1_total_cost"])
    metro_vol = num(balanced["metro_volume"])
    direct_vol = num(balanced["direct_volume"])
    carbon = f"{float(balanced['total_carbon_tCO2']):.4f}"
    road_drop = pct(1 - float(balanced["road_cost"]) / BASELINE_COST)
    direct_drop = pct(1 - float(balanced["direct_volume"]) / BASELINE_DEMAND)
    carbon_drop = pct(1 - float(balanced["total_carbon_tCO2"]) / BASELINE_CARBON)
    cost_change = pct(float(balanced["Z1_total_cost"]) / BASELINE_COST - 1)

    replacements = {
        3: f"针对中心城区地铁物流转运点候选集合构建与启用分配决策衔接不足的问题，本文以成都市武侯区行政边界内全部地铁换乘站为候选输入，构建“换乘站候选集构建—多目标选址-分配优化—案例验证”的研究框架。首先，依据物理站点去重和线路数判定规则形成15座候选转运点集合，并采用 CRITIC-TOPSIS 对站点适宜性进行排序解释；其次，建立以综合成本最小化和地铁服务比例最大化为目标的选址-分配模型，将转运点启用、末端需求分配、运输方式选择、站点容量、线路能力和换乘次数等因素纳入统一框架；最后，采用 QVNS-NSGA-II 求解 Pareto 解集，并通过 TOPSIS 选择代表性折中方案。结果表明，均衡折中方案启用{enabled_names}4座转运点，综合成本为{cost}元/日，地铁服务比例为{service}，地铁联运量为{metro_vol}件/日，总碳排放为{carbon} tCO2/日。与全公路直送基准方案相比，该方案综合成本增加{cost_change}，但公路运输成本下降{road_drop}，公路直送量下降{direct_drop}，总碳排放下降{carbon_drop}。研究表明，将换乘站作为候选转运点输入有助于突出轨道网络跨线组织与货流中转优势，并为中心城区地铁物流节点布局提供可解释的量化依据。",
        5: f"To address the weak connection between candidate-set construction and activation-allocation decisions for metro logistics transfer points, this study uses all 15 metro transfer stations within Wuhou District, Chengdu, as candidate inputs and develops a transfer-station candidate construction, multi-objective location-allocation optimization, and case validation framework. Physical stations are deduplicated and classified by the number of serving lines, while CRITIC-TOPSIS is retained for suitability ranking and interpretation rather than for truncating the candidate set by ranking. A bi-objective model is then formulated to minimize system cost and maximize the metro service ratio under station capacity, line capacity, transfer, and last-mile service constraints. QVNS-NSGA-II is used to generate Pareto solutions, and TOPSIS is adopted to select representative schemes. The balanced compromise solution activates South Railway Station, Gaoshengqiao, Wuqing South Road, and Jitouqiao, with a total cost of {cost} yuan/day, a metro service ratio of {service}, metro volume of {metro_vol} parcels/day, and carbon emissions of {carbon} tCO2/day. Compared with the all-road benchmark, total cost increases by {cost_change}, while road transport cost, direct road volume, and carbon emissions decrease by {road_drop}, {direct_drop}, and {carbon_drop}, respectively.",
        15: "基于上述问题，本文面向中心城区地铁物流配送场景，构建换乘站候选集与多目标选址-分配优化方法。研究目标包括：在明确行政边界和换乘属性判定规则的基础上构建候选转运点集合；在候选集合基础上确定转运点启用方案、末端需求分配关系和运输方式选择；在综合成本、地铁服务比例及运营约束之间形成可解释的权衡方案。具体而言，本文将武侯区内物理去重后的全部换乘站作为小论文算例输入，并保留 CRITIC-TOPSIS 作为站点适宜性排序和结果解释工具，使候选集构建与后续选址-分配模型之间保持清晰边界。",
        16: "本文的主要贡献体现在以下三个方面：一是提出面向地铁物流的换乘站候选集构建口径，将线路数不少于2条的站点纳入候选集合，并区分换乘型站点与枢纽型站点，突出换乘站在跨线组织、货流中转和线路衔接中的作用；二是建立同时刻画转运点启用、需求分配和运输方式选择的双目标选址-分配模型，并纳入站点容量、线路能力、换乘次数和末端接驳距离等轨道交通运营约束；三是将 QVNS-NSGA-II 与 TOPSIS 结合，用于生成和筛选不同成本—服务偏好下的代表性方案，并通过全公路基准和多算法对比验证方案的工程意义和求解可靠性。",
        17: "本文其余部分安排如下：第2节描述地铁物流配送系统边界、问题特征和研究框架；第3节介绍换乘站候选集构建规则和 CRITIC-TOPSIS 适宜性评价方法；第4节构建多目标选址-分配模型并说明 QVNS-NSGA-II 求解算法；第5节以成都市武侯区全部换乘站输入为例开展案例分析；第6节总结研究结论、规划启示和后续研究方向。",
        27: "（3）候选转运点来自武侯区行政边界内全部地铁换乘站集合。物理站点去重后，线路数不少于2条的站点进入候选集合；线路数为2的站点记为TR，线路数不少于3的站点记为HB。启用后的转运点需满足站点容量、最低负载率、最高负载率和相邻转运点间距等约束。",
        31: "本文采用“换乘站候选集构建—多目标选址-分配优化—案例验证”的研究框架。首先，在研究区行政边界内识别全部地铁换乘站，按物理站点去重和线路数规则形成候选集合，并采用 CRITIC-TOPSIS 对站点适宜性进行排序解释。其次，以换乘站候选集合为输入，建立综合成本最小化和地铁服务比例最大化的双目标选址-分配模型，联合刻画转运点启用、需求分配、运输方式选择、站点容量、线路能力、换乘次数和末端接驳半径等约束。最后，利用 QVNS-NSGA-II 输出 Pareto 解集，并通过 TOPSIS 选择不同偏好下的代表性方案。",
        32: "该框架的核心在于将“候选集合构建”与“选址-分配优化”衔接起来。前者保证模型输入具有明确的轨道换乘属性和工程判定规则，后者则在多约束条件下检验候选点组合的系统运行效果，从而避免仅凭站点综合排名直接给出转运点布局方案。",
        52: "在小论文算例中，候选转运点集合不再采用 CRITIC-TOPSIS 排名截取口径，而是限定为武侯区行政边界内全部地铁换乘站。具体判定规则为：对物理站点进行去重，若站点服务线路数不少于2条，则纳入候选集合；线路数为2的站点记为TR换乘型站点，线路数不少于3的站点记为HB枢纽型站点。按该口径，研究区共形成15座候选转运点，包括省体育馆、倪家桥、火车南站、孵化园、世纪城、高升桥、太平园、武青南路、九兴大道、神仙树、锦城大道、高朋大道、三元、华兴和机投桥。",
        53: "CRITIC-TOPSIS 在本文中主要用于解释换乘站适宜性差异，而不再承担截取候选站点数量的筛选功能。最终启用哪些转运点及其服务哪些末端需求点，仍需由选址-分配模型在站点容量、启用成本、线路货运能力、末端服务半径、转运点间距和需求分配关系等约束下进一步确定。",
        54: "表2 武侯区换乘站候选转运点属性表\nTable 2 Attributes of candidate transfer stations in Wuhou District",
        100: "在需求规模设定方面，基准需求情景下研究区日均总需求量为66970.02件/日。模型采用武侯区行政边界内全部15座地铁换乘站作为候选转运点集合，其中12座为TR换乘型站点、3座为HB枢纽型站点，不再包含普通站输入。最大末端接驳距离设为4 km，转运点最小间距设为2 km，每条地铁线路最多启用3个转运点，物流园区接入站至转运点路径最大换乘次数不超过2次，每条线路平峰期可用货运班次为5趟。",
        102: "图2 武侯区换乘站候选转运点空间分布\nFig.2 Spatial distribution of candidate transfer stations in Wuhou District",
        106: "案例中的候选转运点来自武侯区行政边界内全部地铁换乘站，而非按 CRITIC-TOPSIS 排名截取的候选集合。候选集合共15座，包括12座TR换乘型站点和3座HB枢纽型站点。模型根据站点类型映射处理能力和固定启用成本：TR站容量为7000件/平峰作业周期，固定启用成本为12000元/日；HB站容量为10000件/平峰作业周期，固定启用成本为16000元/日。由于普通站不进入本算例，普通站容量和成本参数不参与本次小论文重算。",
        107: "需要强调的是，换乘站候选集构建只规定模型输入边界，并不意味着全部换乘站都会被启用。最终启用结果还受到末端需求分布、地铁线路可达性、换乘次数、站点容量、线路班次能力、转运点间距和综合成本等因素共同影响。因此，案例分析的重点不是重复展示站点评价排序，而是检验全部换乘站输入在选址-分配模型中的组合效果。",
        110: f"三类偏好方案表现出清晰的权衡关系。成本偏好方案的综合成本为{money(prefs.iloc[0]['Z1_total_cost'])}元/日，地铁服务比例为{pct(prefs.iloc[0]['Z2_metro_service_ratio'])}，启用{intish(prefs.iloc[0]['num_open'])}个转运点；均衡偏好方案的综合成本为{money(prefs.iloc[1]['Z1_total_cost'])}元/日，地铁服务比例为{pct(prefs.iloc[1]['Z2_metro_service_ratio'])}，启用{intish(prefs.iloc[1]['num_open'])}个转运点；服务偏好方案的综合成本为{money(prefs.iloc[2]['Z1_total_cost'])}元/日，地铁服务比例为{pct(prefs.iloc[2]['Z2_metro_service_ratio'])}，启用{intish(prefs.iloc[2]['num_open'])}个转运点。随着地铁服务比例提高，公路直送量由{num(prefs.iloc[0]['direct_volume'])}件/日降至{num(prefs.iloc[2]['direct_volume'])}件/日，但系统需要承担更高的启用、接驳和轨道组织成本。",
        113: "图3 15座换乘站输入下的 Pareto 前沿与均衡折中方案\nFig.3 Pareto front and balanced compromise solution under the 15-transfer-station input",
        117: f"在均衡偏好下，TOPSIS 选取的代表性折中解编号为{intish(balanced['solution_id'])}，启用转运点为{enabled_names}。该方案启用1个HB枢纽型站点和3个TR换乘型站点，固定启用成本合计{intish(balanced['fixed_cost_total'])}元/日，地铁联运量为{metro_vol}件/日，公路直送量为{direct_vol}件/日，地铁服务比例达到{service}。方案平均末端接驳距离为{num(balanced['avg_feeder_distance'])} km，满足4 km服务半径约束；占用平峰期地铁货运班次{intish(balanced['used_train_trips_total'])}趟，最大站点负载率为{pct(balanced['max_load_rate'])}，最大线路装载率为{pct(balanced['max_line_load_rate'])}，说明该方案在站点容量和线路能力约束下保持可行。",
        125: f"为检验地铁联运方案的工程意义，本文将均衡折中方案与全公路直送基准方案进行对比。全公路直送方案的综合成本为204130.69元/日，地铁服务比例为0，全部66970.02件/日需求由公路完成；均衡折中方案的综合成本为{cost}元/日，较全公路基准增加{cost_change}，但公路运输成本由204130.69元/日下降至{money(balanced['road_cost'])}元/日，降幅为{road_drop}；公路直送量下降{direct_drop}，总碳排放由0.5102 tCO2/日下降至{carbon} tCO2/日，降幅为{carbon_drop}。这表明换乘站限定输入下的地铁物流方案更主要体现为道路运输替代和低碳减排收益，而不是单纯追求总成本最低。",
        133: f"为验证 QVNS-NSGA-II 对15座换乘站输入算例的求解有效性，本文将其与 NSGA-II、MOPSO 和 MOEA/D 进行同口径对比。结果显示，QVNS-NSGA-II 的 HV 均值为{float(algo.loc[algo['algorithm']=='QVNS-NSGA-II','HV_mean'].iloc[0]):.4f}，高于 NSGA-II、MOPSO 和 MOEA/D；IGD 均值为{float(algo.loc[algo['algorithm']=='QVNS-NSGA-II','IGD_mean'].iloc[0]):.4f}，低于其他对比算法；Spacing 均值为{float(algo.loc[algo['algorithm']=='QVNS-NSGA-II','Spacing_mean'].iloc[0]):.4f}，表明解集分布更均匀。该结果说明，在新的换乘站候选输入下，QVNS-NSGA-II 仍能提供覆盖质量较高且具有工程解释性的 Pareto 解集。",
        138: "本文面向中心城区地铁物流配送场景，提出了“换乘站候选集构建—多目标选址-分配优化—案例验证”的研究框架。首先，依据行政边界、物理站点去重和线路数判定规则构建地铁换乘站候选集合，并采用 CRITIC-TOPSIS 方法进行站点适宜性排序解释；其次，建立综合成本最小化与地铁服务比例最大化的双目标选址-分配模型，联合考虑转运点启用、需求分配、运输方式选择、站点容量、线路能力、换乘次数和末端接驳距离等约束；最后，采用 QVNS-NSGA-II 求解 Pareto 解集，并通过 TOPSIS 选择代表性折中方案。",
        139: f"成都武侯区案例结果表明，在全部15座换乘站输入下，均衡折中方案启用{enabled_names}4个地铁物流转运点，地铁服务比例达到{service}，地铁联运量为{metro_vol}件/日，平均末端接驳距离为{num(balanced['avg_feeder_distance'])} km，且站点容量、线路班次和服务半径等约束均满足要求。与全公路直送基准方案相比，该方案综合成本增加{cost_change}，但公路运输成本下降{road_drop}，公路直送量下降{direct_drop}，总碳排放下降{carbon_drop}。这说明地铁物流在中心城区配送中更适合作为道路配送的协同补充方式，其价值主要体现在道路运输替代、轨道资源利用和低碳配送潜力方面。",
        140: "从规划启示看，中心城区地铁物流转运点选择不宜仅依据站点排名或单一需求强度判断，而应综合考虑换乘站跨线衔接能力、线路可达性、站点容量、周边需求潜力和末端接驳范围。以全部换乘站作为候选输入，可以突出轨道网络组织优势并减少普通站货运化假设带来的不确定性；选址-分配优化则能够在多约束条件下形成可解释、可比较的布局方案。",
    }
    for idx, text in replacements.items():
        set_paragraph_text(doc.paragraphs[idx], text)

    replace_image(doc.paragraphs[33], figs["fig1"], 5.95)
    replace_image(doc.paragraphs[101], figs["fig2"], 5.95)
    replace_image(doc.paragraphs[112], figs["fig3"], 5.95)
    replace_image(doc.paragraphs[120], figs["fig4"], 5.95)
    replace_image(doc.paragraphs[128], figs["fig5"], 5.95)

    fill_table(doc.tables[1], table_candidate_rows(candidates))
    fill_table(doc.tables[3], table_parameter_rows())
    fill_table(doc.tables[5], table_data_source_rows())
    fill_table(doc.tables[6], table_preference_rows(prefs))
    fill_table(doc.tables[7], table_load_rows(station_load))
    fill_table(doc.tables[8], table_baseline_rows(balanced))
    fill_table(doc.tables[9], table_algorithm_rows(algo))


def markdown_table(rows: list[list[str]]) -> str:
    widths = [max(len(str(row[i])) for row in rows) for i in range(len(rows[0]))]
    out = []
    out.append("| " + " | ".join(str(v).ljust(widths[i]) for i, v in enumerate(rows[0])) + " |")
    out.append("| " + " | ".join("-" * widths[i] for i in range(len(widths))) + " |")
    for row in rows[1:]:
        out.append("| " + " | ".join(str(v).ljust(widths[i]) for i, v in enumerate(row)) + " |")
    return "\n".join(out)


def write_note(candidates: pd.DataFrame, prefs: pd.DataFrame, balanced: pd.Series, station_load: pd.DataFrame, algo: pd.DataFrame) -> None:
    enabled_names = "、".join(station_load.loc[pd.to_numeric(station_load["是否启用"], errors="coerce") > 0.5, "地铁站名称"].tolist())
    lines = [
        "# 小论文换乘站输入修订说明",
        "",
        "## 修订口径",
        "",
        "- 小论文已从“CRITIC-TOPSIS 前17个候选站点”改为“武侯区行政边界内全部地铁换乘站”输入。",
        "- 判定规则：物理站点去重，线路数不少于2条；线路数为2记为TR，线路数不少于3记为HB。",
        f"- 候选站规模：{len(candidates)}座，其中TR {int((candidates['station_type']=='TR').sum())}座，HB {int((candidates['station_type']=='HB').sum())}座；普通站不进入小论文算例输入。",
        "",
        "## 均衡折中方案",
        "",
        f"- solution_id：{intish(balanced['solution_id'])}",
        f"- 启用站点：{enabled_names}",
        f"- 综合成本：{money(balanced['Z1_total_cost'])} 元/日",
        f"- 地铁服务比例：{pct(balanced['Z2_metro_service_ratio'])}",
        f"- 地铁联运量：{num(balanced['metro_volume'])} 件/日",
        f"- 公路直送量：{num(balanced['direct_volume'])} 件/日",
        f"- 总碳排放：{float(balanced['total_carbon_tCO2']):.4f} tCO2/日",
        f"- 相对全公路基准：综合成本变化 {pct(float(balanced['Z1_total_cost']) / BASELINE_COST - 1)}，公路运输成本变化 {pct(float(balanced['road_cost']) / BASELINE_COST - 1)}，公路直送量变化 {pct(float(balanced['direct_volume']) / BASELINE_DEMAND - 1)}，碳排放变化 {pct(float(balanced['total_carbon_tCO2']) / BASELINE_CARBON - 1)}。",
        "",
        "## 三类偏好方案",
        "",
        markdown_table(table_preference_rows(prefs)),
        "",
        "## 算法对比",
        "",
        markdown_table(table_algorithm_rows(algo)),
        "",
        "## 输出文件",
        "",
        f"- 修订版Word：`{OUT_DOCX}`",
        f"- 本说明：`{OUT_NOTE}`",
        f"- 新图表资产目录：`{ASSET_DIR}`",
    ]
    OUT_NOTE.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    ensure_paths()
    candidates = read_csv(CANDIDATE_PATH)
    pareto = read_csv(PARETO_PATH)
    station_load = read_csv(STATION_LOAD_PATH)
    assignments = read_csv(ASSIGNMENTS_PATH)
    algo = read_csv(ALGO_SUMMARY_PATH)
    prefs = preference_solutions(pareto)
    balanced = pareto[pareto["solution_id"].astype(int) == int(prefs.iloc[1]["solution_id"])].iloc[0]

    figs = {
        "fig1": save_flowchart(),
        "fig2": save_candidate_map(candidates),
        "fig3": save_pareto_plot(pareto, prefs),
        "fig4": save_compromise_map(candidates, station_load, assignments),
        "fig5": save_baseline_chart(balanced),
    }

    doc = Document(str(SOURCE_DOCX))
    update_document(doc, candidates, prefs, balanced, station_load, algo, figs)
    doc.save(str(OUT_DOCX))
    write_note(candidates, prefs, balanced, station_load, algo)

    print(f"wrote_docx={OUT_DOCX}")
    print(f"wrote_note={OUT_NOTE}")
    for name, path in figs.items():
        print(f"{name}={path}")


if __name__ == "__main__":
    main()
