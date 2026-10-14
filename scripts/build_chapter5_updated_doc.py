# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
RESULT = ROOT / "results" / "four_algorithm_iteration_p100_g300_runs10"
SRC_NAME = "第五章_成都市武侯区地铁物流配送转运点布局优化研究_优化版_图已同步.docx"


def set_run_font(run, size=10.5, bold=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def format_paragraph(paragraph, first_line=True):
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    for run in paragraph.runs:
        set_run_font(run)


def cell_text(cell, text, bold=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=9, bold=bold)


def main() -> int:
    src = DOCS / SRC_NAME
    doc = Document(src)
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    start = end = None
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("5.3.4"):
            start = index
        if start is not None and index > start and text.startswith("5.4"):
            end = index
            break
    if start is None or end is None:
        raise RuntimeError("Cannot locate section 5.3.4")

    for paragraph in doc.paragraphs[start:end]:
        paragraph._element.getparent().remove(paragraph._element)

    anchor_paragraph = doc.paragraphs[start]

    def insert_paragraph(text="", style=None):
        paragraph = anchor_paragraph.insert_paragraph_before(text)
        if style:
            paragraph.style = style
        return paragraph

    def add_text(text, style=None, first_line=True):
        paragraph = insert_paragraph(text, style)
        format_paragraph(paragraph, first_line=first_line)
        return paragraph

    def add_caption(text):
        paragraph = add_text(text, first_line=False)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=10)
        return paragraph

    def add_image(filename, width):
        paragraph = insert_paragraph("")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(RESULT / filename), width=Inches(width))

    def add_performance_table():
        perf = pd.read_csv(RESULT / "paper_table_algorithm_performance.csv", encoding="utf-8-sig")
        table = doc.add_table(rows=perf.shape[0] + 1, cols=perf.shape[1])
        anchor_paragraph._p.addprevious(table._tbl)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for col_index, col_name in enumerate(perf.columns):
            cell_text(table.cell(0, col_index), col_name, bold=True)
        for row_index, (_, row) in enumerate(perf.iterrows(), start=1):
            for col_index, col_name in enumerate(perf.columns):
                cell_text(table.cell(row_index, col_index), row[col_name])

    add_text("5.3.4 QVNS-NSGA-II 与典型多目标优化算法的性能对比", style="Heading 3", first_line=False)
    add_text("为验证 QVNS-NSGA-II 在地铁物流配送转运点布局优化问题中的有效性，本文将其与 NSGA-II、MOPSO 和 SPEA2 三类典型多目标优化算法进行对比。四类算法均在相同的候选转运点、需求矩阵、成本参数和约束条件下运行，种群规模设置为 100，最大迭代代数设置为 300，并分别开展 10 次独立实验。评价指标包括超体积指标（HV）、反世代距离（IGD）、Spacing、Spread、非重复解数和运行时间。其中，HV 越大表示解集的收敛性和覆盖能力越好，IGD 越小表示解集越接近参考前沿，Spacing 和 Spread 用于刻画解集分布均匀性与扩展性。")
    add_text("从 Pareto 前沿分布看，QVNS-NSGA-II 在综合成本—地铁服务比例目标空间中形成了更靠近左上方的前沿区域，说明其能够在较低综合成本下获得较高地铁服务比例。相比之下，NSGA-II 的前沿覆盖范围略窄，MOPSO 与 SPEA2 在部分高服务比例区域的探索能力不足，整体前沿质量弱于 QVNS-NSGA-II。")
    add_image("paper_fig_pareto_front_cost_service.png", 5.8)
    add_caption("图 5-3 典型多目标优化算法 Pareto 前沿对比")
    add_text("为进一步观察算法迭代过程，本文统计各算法在迭代过程中的 Pareto 前沿平均总成本、Pareto 前沿平均地铁服务比例、HV 和成本改进率。由图 5-4 可见，QVNS-NSGA-II 在迭代过程中能够持续改善前沿质量，并在 HV 与成本改进率方面保持较优表现，说明 Q-learning 动作选择、变邻域搜索和外部档案机制能够增强算法在离散选址—分配空间中的搜索效率。")
    add_image("paper_fig_algorithm_convergence_grid.png", 6.2)
    add_caption("图 5-4 四种算法收敛过程对比")
    add_text("图 5-5 展示了 QVNS-NSGA-II 单算法迭代过程中 Pareto 前沿平均总成本和平均地铁服务比例的变化。可以看出，算法在前期快速调整解集结构，随后进入相对稳定的局部改进阶段，体现出增强型 NSGA-II 在成本控制和服务比例提升之间的动态权衡能力。")
    add_image("paper_fig_iteration_effect_qvns_nsga2.png", 5.8)
    add_caption("图 5-5 QVNS-NSGA-II 迭代效果")
    add_text("为比较不同算法的整体求解性能，本文进一步统计 10 次独立运行的平均指标，结果如表 5-10 和图 5-6 所示。QVNS-NSGA-II 的 HV 均值为 0.8499，高于 NSGA-II 的 0.7841、MOPSO 的 0.5821 和 SPEA2 的 0.6676；IGD 均值为 0.1438，低于 NSGA-II 的 0.1720、MOPSO 的 0.3246 和 SPEA2 的 0.2606，说明 QVNS-NSGA-II 在逼近参考前沿和维持前沿覆盖能力方面具有更优表现。")
    add_caption("表 5-10 典型多目标优化算法性能对比")
    add_performance_table()
    add_image("paper_fig_algorithm_performance_metrics.png", 6.2)
    add_caption("图 5-6 四种算法性能指标对比")
    add_text("从运行效率看，QVNS-NSGA-II 的平均运行时间为 93.12 s，低于 NSGA-II、MOPSO 和 SPEA2，表明改进算法并未因引入 Q-learning、VNS 和外部 archive 机制而造成不可接受的计算负担。相反，基于反馈的动作选择和自适应局部搜索机制提高了有效搜索比例，使算法在保持较高解集质量的同时具备较好的求解效率。")
    add_text("图 5-7 给出了四类算法 10 次独立运行下的稳定性分布。QVNS-NSGA-II 在 HV 和 IGD 指标上整体分布更优，且运行时间分布相对集中，说明其在不同随机初始条件下均能保持较稳定的前沿质量。相比之下，MOPSO 与 SPEA2 在部分指标上的离散程度更高，反映出其搜索过程受初始种群和随机扰动影响更明显。")
    add_image("paper_fig_algorithm_stability_distribution.png", 6.2)
    add_caption("图 5-7 多次运行稳定性分布")
    add_text("综上，QVNS-NSGA-II 的性能优势主要来自三个方面。第一，Q-learning 能够根据搜索反馈动态调整进化动作，提高不同搜索阶段的参数适应性；第二，自适应 VNS 在前沿改善停滞时触发局部邻域扰动，有助于跳出局部最优并改善解集质量；第三，外部档案与边界解保留机制增强了低成本端和高服务端解的保留能力，使算法能够更稳定地覆盖成本—服务比例权衡前沿。因此，QVNS-NSGA-II 更适用于本文所研究的带复杂约束的地铁物流配送转运点选址—分配双目标优化问题。")

    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "MOEA/D" in text or "20 次" in text or "20次" in text:
            if "本章以成都市武侯区" in text and "QVNS-NSGA-II" in text:
                paragraph.text = "本章以成都市武侯区为案例区域，基于地铁站点、顺丰末端站点、物流园区、路径距离矩阵和运营参数等数据，开展了地铁物流配送转运点布局优化研究。首先，利用 CRITIC-TOPSIS 评价结果筛选得到 17 个备选地铁物流配送转运点，并根据站点类型设定处理能力和固定启用成本。其次，构建无预算约束下的地铁物流选址—分配双目标模型，将综合成本最小化和地铁服务比例最大化作为优化目标，并纳入末端服务距离、换乘次数、站点负载率、线路启用数量和平峰期班次容量等约束。再次，采用 QVNS-NSGA-II 算法进行求解，得到覆盖低成本端与高服务端的 Pareto 解集。最后，与 NSGA-II、MOPSO 和 SPEA2 的 10 次独立实验对比表明，QVNS-NSGA-II 在 HV、IGD、运行时间和多次运行稳定性方面具有综合优势，能够在可接受计算时间内获得质量更高、分布更稳定的 Pareto 前沿。敏感性分析表明，需求增长会提高系统成本并压缩地铁服务比例，成本参数同比例变化主要改变成本规模而不改变布局结构，货源比例变化则会显著影响成本水平和服务覆盖方向。整体来看，地铁物流转运点布局应优先关注具备多线连接能力、靠近需求密集区且负载可控的枢纽及换乘站点，同时需要结合平峰期班次组织能力和物流源点结构进行动态调整。"
            else:
                paragraph.text = text.replace("MOEA/D", "SPEA2").replace("20 次", "10 次").replace("20次", "10次")
            format_paragraph(paragraph)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_run_font(run)

    tmp_docx = DOCS / "chapter5_output_tmp.docx"
    tmp_md = DOCS / "chapter5_output_tmp.md"
    doc.save(tmp_docx)

    md = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            level = 2 if "2" in paragraph.style.name else 3 if "3" in paragraph.style.name else 1
            md.append("#" * level + " " + text)
        else:
            md.append(text)
    md.append("\n## 表 5-10 典型多目标优化算法性能对比\n")
    perf = pd.read_csv(RESULT / "paper_table_algorithm_performance.csv", encoding="utf-8-sig")
    md.append(perf.to_markdown(index=False))
    md.append("\n## 本次更新使用的图表文件\n")
    for name in [
        "paper_fig_pareto_front_cost_service.png",
        "paper_fig_algorithm_convergence_grid.png",
        "paper_fig_iteration_effect_qvns_nsga2.png",
        "paper_fig_algorithm_performance_metrics.png",
        "paper_fig_algorithm_stability_distribution.png",
    ]:
        md.append(f"- {name}")
    tmp_md.write_text("\n\n".join(md), encoding="utf-8")
    print(tmp_docx)
    print(tmp_md)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
