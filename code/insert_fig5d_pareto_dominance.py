# -*- coding: utf-8 -*-
"""Generate Figure 5-D and insert it into the Chapter 5 DOCX.

The figure uses the existing Chapter 5 algorithm comparison Pareto points and
emphasizes the geometric dominance relation in the cost-service objective
space: lower cost and higher metro service ratio are preferred.
"""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
RESULT_DIR = ROOT / "results" / "chapter5_algorithm_comparison"
POINTS_CSV = RESULT_DIR / "pareto_cost_service_points.csv"
FIG_DIR = ROOT / "docs" / "第五章_标准版图片"
FIG_PNG = FIG_DIR / "fig5_D_pareto_dominance_front.png"
FIG_SVG = FIG_DIR / "fig5_D_pareto_dominance_front.svg"
DOC_DIR = ROOT / "results" / "5.20"
DOCX_SUFFIX = "_任务包三Pareto前沿支配优化版.docx"

ALG_ORDER = ["QVNS-NSGA2", "NSGA2", "MOPSO", "SPEA2"]
ALG_LABEL = {
    "QVNS-NSGA2": "QVNS-NSGA-Ⅱ",
    "NSGA2": "NSGA-Ⅱ",
    "MOPSO": "MOPSO",
    "SPEA2": "SPEA2",
}
PALETTE = {
    "QVNS-NSGA2": "#B23A2F",
    "NSGA2": "#3E7CB1",
    "MOPSO": "#4C956C",
    "SPEA2": "#7B6BA8",
}


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [
                "Microsoft YaHei",
                "SimHei",
                "SimSun",
                "Arial Unicode MS",
                "DejaVu Sans",
            ],
            "axes.unicode_minus": False,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "figure.dpi": 180,
            "savefig.dpi": 600,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )


def nondominated_front(df: pd.DataFrame) -> pd.DataFrame:
    """Return the cost-min/service-max non-dominated envelope."""
    front = (
        df[["cost_10k", "service_pct"]]
        .round({"cost_10k": 3, "service_pct": 3})
        .drop_duplicates()
        .sort_values(["cost_10k", "service_pct"])
        .reset_index(drop=True)
    )
    keep: list[int] = []
    best_service = -np.inf
    for idx, row in front.iterrows():
        service = float(row["service_pct"])
        if service >= best_service - 1e-9:
            keep.append(idx)
            best_service = max(best_service, service)
    return front.loc[keep].reset_index(drop=True)


def load_points() -> pd.DataFrame:
    points = pd.read_csv(POINTS_CSV, encoding="utf-8-sig")
    points = points[points["algorithm"].isin(ALG_ORDER)].copy()
    points["cost_10k"] = points["Z1_total_cost_10k"].astype(float)
    points["service_pct"] = points["Z2_metro_service_percent"].astype(float)
    return points


def generate_figure() -> None:
    configure_matplotlib()
    points = load_points()
    FIG_DIR.mkdir(parents=True, exist_ok=True)

    fig, ax = plt.subplots(figsize=(8.4, 5.6))
    for alg in ALG_ORDER:
        subset = points[points["algorithm"] == alg].copy()
        unique = (
            subset[["cost_10k", "service_pct"]]
            .round({"cost_10k": 3, "service_pct": 3})
            .drop_duplicates()
        )
        is_qvns = alg == "QVNS-NSGA2"
        ax.scatter(
            unique["cost_10k"],
            unique["service_pct"],
            s=15 if is_qvns else 11,
            color=PALETTE[alg],
            alpha=0.24 if is_qvns else 0.16,
            edgecolors="none",
            label=f"{ALG_LABEL[alg]} 解集",
        )

        front = nondominated_front(subset)
        ax.plot(
            front["cost_10k"],
            front["service_pct"],
            color=PALETTE[alg],
            lw=2.8 if is_qvns else 1.6,
            alpha=0.96 if is_qvns else 0.72,
            label=f"{ALG_LABEL[alg]} 前沿",
        )
        ax.scatter(
            front["cost_10k"],
            front["service_pct"],
            s=34 if is_qvns else 22,
            color=PALETTE[alg],
            edgecolors="white",
            linewidths=0.55,
            zorder=4,
        )

    ax.annotate(
        "成本降低",
        xy=(19.35, 51.2),
        xytext=(21.1, 51.2),
        arrowprops=dict(arrowstyle="->", lw=1.0, color="#444444"),
        color="#444444",
        ha="center",
        va="center",
        fontsize=10,
    )
    ax.annotate(
        "服务提升",
        xy=(19.55, 53.0),
        xytext=(19.55, 43.5),
        arrowprops=dict(arrowstyle="->", lw=1.0, color="#444444"),
        color="#444444",
        ha="center",
        va="center",
        rotation=90,
        fontsize=10,
    )
    ax.text(
        19.72,
        48.2,
        "Pareto 优势方向：左上方",
        color="#333333",
        fontsize=10,
        bbox=dict(boxstyle="round,pad=0.25", facecolor="white", edgecolor="#CCCCCC", alpha=0.88),
    )

    ax.set_title("多算法非劣解集前沿面几何支配对比", fontsize=14, pad=12)
    ax.set_xlabel("综合成本 Z1（万元/日）", fontsize=11)
    ax.set_ylabel("地铁服务比例 Z2（%）", fontsize=11)
    ax.set_xlim(19.2, 25.4)
    ax.set_ylim(-1.5, 55.5)
    ax.grid(True, linestyle="--", linewidth=0.55, alpha=0.30)

    handles, labels = ax.get_legend_handles_labels()
    # Keep only front lines in the legend to reduce clutter.
    front_handles = [h for h, label in zip(handles, labels) if label.endswith("前沿")]
    front_labels = [label.replace(" 前沿", "") for label in labels if label.endswith("前沿")]
    ax.legend(
        front_handles,
        front_labels,
        loc="upper center",
        bbox_to_anchor=(0.5, -0.13),
        ncol=4,
        frameon=False,
        handlelength=2.0,
        columnspacing=1.3,
    )
    fig.subplots_adjust(bottom=0.22)
    fig.savefig(FIG_PNG, bbox_inches="tight")
    fig.savefig(FIG_SVG, bbox_inches="tight")
    plt.close(fig)


def latest_chapter5_docx() -> Path:
    candidates = [
        p
        for p in DOC_DIR.glob("*.docx")
        if not p.name.startswith("~$")
        and "第五章" in p.name
        and "任务包三Pareto前沿支配优化版" not in p.name
    ]
    if not candidates:
        raise FileNotFoundError("未找到第五章 Word 文档。")
    return max(candidates, key=lambda p: p.stat().st_mtime)


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "宋体")


def qn(tag: str) -> str:
    prefix, tagroot = tag.split(":")
    nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return f"{{{nsmap[prefix]}}}{tagroot}"


def paragraph_before(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        set_run_font(run)
    return new_para


def insert_picture_before(paragraph, image_path: Path, width_cm: float = 15.5):
    p = paragraph_before(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    return p


def insert_fig5d_into_docx() -> Path:
    source = latest_chapter5_docx()
    target = source.with_name(source.stem + DOCX_SUFFIX)
    shutil.copy2(source, target)
    doc = Document(target)

    if any("图5-D 多算法非劣解集前沿面几何支配对比" in p.text for p in doc.paragraphs):
        doc.save(target)
        return target

    anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("图5-7 QVNS-NSGA-II 与典型多目标优化算法性能对比"):
            anchor = paragraph
            break
    if anchor is None:
        raise ValueError("未找到图5-7标题段，无法定位插入位置。")

    intro_text = (
        "在统计指标比较之前，为直观呈现不同算法在双目标空间中的支配关系，本文进一步绘制多算法非劣解集前沿面几何支配对比图。"
        "图 5-D 以综合成本 $Z_1$ 为横轴、地铁服务比例 $Z_2$ 为纵轴，其中成本越低、服务比例越高表示方案越优，因此目标空间左上方代表 Pareto 优势方向。"
        "从前沿包络形态看，QVNS-NSGA-Ⅱ 在低成本端和高服务端均保持较完整的非劣解覆盖，相较于 NSGA-Ⅱ、MOPSO 和 SPEA2，其前沿整体更靠近左上方，说明改进算法不仅在均值指标上表现较好，也能在物理可解释的成本—服务权衡空间中提供更具支配性的候选方案集。"
    )
    intro = paragraph_before(anchor, intro_text, style="Normal")
    intro.paragraph_format.first_line_indent = Cm(0.74)
    for run in intro.runs:
        set_run_font(run)

    insert_picture_before(anchor, FIG_PNG)
    caption = paragraph_before(anchor, "图5-D 多算法非劣解集前沿面几何支配对比", style="Normal")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        set_run_font(run, size=10.5)

    analysis_text = (
        "该图与图 5-7 的统计指标形成互补：图 5-D 侧重展示非劣解集在目标空间中的几何位置和支配方向，"
        "图 5-7 则从 HV、IGD、Spacing、Spread、唯一解数量和运行时间等指标对算法性能进行量化检验。"
        "二者结合能够同时说明 QVNS-NSGA-Ⅱ 的前沿覆盖优势、收敛质量和计算代价特征。"
    )
    analysis = paragraph_before(anchor, analysis_text, style="Normal")
    analysis.paragraph_format.first_line_indent = Cm(0.74)
    for run in analysis.runs:
        set_run_font(run)

    doc.save(target)
    return target


def main() -> int:
    generate_figure()
    output_docx = insert_fig5d_into_docx()
    print(f"Figure saved: {FIG_PNG}")
    print(f"Document saved: {output_docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
