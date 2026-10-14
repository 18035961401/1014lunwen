from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SRC_MD = ROOT / "第五章_成都市武侯区地铁物流配送转运点布局优化研究_必须填充最终版.md"
OUT_MD = ROOT / "第五章_成都市武侯区地铁物流配送转运点布局优化研究_成本能力参数补强最终版.md"
OUT_DOCX = ROOT / "第五章_成都市武侯区地铁物流配送转运点布局优化研究_成本能力参数补强最终版.docx"


def insert_after(text: str, anchor: str, addition: str) -> str:
    if addition.strip() in text:
        return text
    if anchor not in text:
        raise RuntimeError("anchor not found")
    return text.replace(anchor, anchor + "\n\n" + addition.strip(), 1)


def update_text(text: str) -> str:
    anchor = (
        "| 每线可用班次数 | 5 | 平峰期可用列车班次 |"
    )
    addition = """
为兑现第四章综合成本函数 $Z_1$ 中各成本子项的参数标定要求，本文进一步给出全链条协同运营成本与能力参数。表 5-X 汇总了公路直送、地铁干线、入网短驳、出网接驳、换乘作业、进出站作业、固定启用成本、班次组织成本以及碳外部成本的取值口径。需要说明的是，表中“件”在模型计算中统一按 SLU 当量理解，即一个顺丰 F3 标准箱折算为一个 SLU；若后续企业提供更精细的箱型或订单结构数据，可在保持模型结构不变的前提下重新标定单位成本。

| 参数类别 | 符号或代码变量 | 参数值 | 单位 | 对应公式（4-1）成本子项或约束作用 | 数据来源或说明 |
|---|---|---:|---|---|---|
| 公路直送单位运输成本 | $c^R$ / road_operating_cost | 0.12 | 元/(SLU·km) | 公路直送运输成本、入网短驳距离相关成本 | 模型全局参数_无预算.csv |
| 地铁干线单位运输成本 | $c^M$ / metro_electric_cost | 0.02 | 元/(SLU·km) | 地铁干线运输距离相关成本 | 模型全局参数_无预算.csv |
| 末端接驳单位运输成本 | $c^F$ / feeder_cost | 0.04 | 元/(SLU·km) | 转运点至顺丰营业网点末端接驳成本 | 模型全局参数_无预算.csv |
| 地铁换乘作业成本 | $c^T$ / transfer_cost | 0.10 | 元/(SLU·次) | 跨线换乘搬运、交接与组织成本 | 模型全局参数_无预算.csv |
| 进出站作业成本 | $c^{IO}$ / inout_cost | 0.20 | 元/SLU | 货物安检、垂直上下线、站内交接成本 | 模型全局参数_无预算.csv |
| SL/TR/HB 固定启用成本 | $F_k$ / fixed_cost_SL/TR/HB | 8000/12000/16000 | 元/日 | 候选转运点启用固定成本 | 模型全局参数_无预算.csv；候选点表按站点类型匹配 |
| 班次组织成本率 | $\\theta_l$ / train_organization_cost_rate | 0.06 | 元/SLU | 平峰期列车参与货运的组织与空间占用成本 | 模型全局参数_无预算.csv |
| 碳价格 | $p^C$ / carbon_price | 0.10 | 元/kgCO2 | 碳排放外部成本折算参数 | 模型全局参数_无预算.csv |
| 公路排放因子 | $e^R$ / road_emission_factor | 0.00030 | kgCO2/(SLU·km) | 公路直送及短驳碳排放成本 | 模型全局参数_无预算.csv |
| 地铁排放因子 | $e^M$ / metro_emission_factor | 0.00003 | kgCO2/(SLU·km) | 地铁干线运输碳排放成本 | 模型全局参数_无预算.csv |
| 末端接驳排放因子 | $e^F$ / feeder_emission_factor | 0.00010 | kgCO2/(SLU·km) | 新能源末端接驳车辆碳排放成本 | 模型全局参数_无预算.csv |
| 每线可用平峰班次数 | $A_l$ / available_train_trips | 5 | 趟/线路 | 地铁线路平峰期可用班次约束 | 模型全局参数_无预算.csv |

表 5-Y 进一步给出线路级列车装载能力参数。单趟列车最大装载能力 $G_l$ 由顺丰 F3 标准箱尺寸、车型编组、车厢可用货运体积比例和三维整数装箱计算得到；班次组织成本按 $\\theta_l G_l$ 计算，仅表示既有平峰客运列车参与货运的增量组织与空间占用成本，并非完整列车运行成本。

| 地铁线路 | 车辆类型 | 编组 | 单趟列车最大装载能力 $G_l$（SLU/趟） | 可用班次数 $A_l$（趟） | 线路总理论能力（SLU） | 班次组织成本率 $\\theta_l$（元/SLU） | 单趟班次组织成本（元/趟） |
|---|---|---:|---:|---:|---:|---:|---:|
| 1号线 | B型车 | 6 | 7560 | 5 | 37800 | 0.06 | 453.60 |
| 3号线 | B型车 | 6 | 7560 | 5 | 37800 | 0.06 | 453.60 |
| 5号线 | A型车 | 8 | 12936 | 5 | 64680 | 0.06 | 776.16 |
| 7号线 | A型车 | 6 | 9702 | 5 | 48510 | 0.06 | 582.12 |
| 8号线 | A型车 | 6 | 9702 | 5 | 48510 | 0.06 | 582.12 |
| 10号线 | A型车 | 6 | 9702 | 5 | 48510 | 0.06 | 582.12 |
| 18号线 | 市域A型车 | 8 | 12936 | 5 | 64680 | 0.06 | 776.16 |

上述两张参数表与表 5-5 中候选站点容量和固定启用成本共同构成公式（4-1）综合成本函数 $Z_1$ 的实证参数输入。其中，时间成本不再折算入综合成本，仅作为服务水平统计指标；公路基础分拣与交接成本视为不同运输方式共同发生的基础成本，不作为差异化成本重复计入。由此，第四章承诺的单位综合运输成本、转运点启用成本、线路班次组织成本、列车装载能力和碳外部成本均在本节完成标定。
"""
    return insert_after(text, anchor, addition)


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(doc: Document, text: str, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, 10.5)


def add_table(doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, 10.5, bold=(i == 0))


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(set(c.replace(":", "").replace("-", "").strip()) == set() for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def markdown_to_docx(md_text: str, out_path: Path):
    doc = Document()
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    lines = md_text.splitlines()
    i = 0
    para_buf = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            add_paragraph(doc, " ".join(x.strip() for x in para_buf if x.strip()))
            para_buf = []

    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            flush_para()
            i += 1
            continue
        if stripped.startswith("|"):
            flush_para()
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
            continue
        if stripped.startswith("#"):
            flush_para()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            p = doc.add_paragraph(style=style)
            r = p.add_run(title)
            set_run_font(r, 14 if level == 1 else 12 if level == 2 else 11, bold=True)
            i += 1
            continue
        para_buf.append(stripped)
        i += 1
    flush_para()
    doc.save(out_path)


def main():
    text = SRC_MD.read_text(encoding="utf-8-sig")
    updated = update_text(text)
    OUT_MD.write_text(updated, encoding="utf-8")
    markdown_to_docx(updated, OUT_DOCX)
    print(OUT_MD)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
