from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt


ROOT = Path(r"D:\trae建模\metro_logistics_qvns_nsga2_package\metro_logistics_qvns_nsga2_package")
TARGET = ROOT / "results" / "5.20" / "第五章_成都市武侯区地铁物流配送转运点布局优化研究_成本能力参数补强最终版.docx"
FALLBACK_TARGET = ROOT / "results" / "5.20" / "第五章_成都市武侯区地铁物流配送转运点布局优化研究_成本能力参数补强最终版_表图核验更新版.docx"
OUT_DIR = ROOT / "results" / "5.20" / "ch5_visual_updates"
OUT_DIR.mkdir(parents=True, exist_ok=True)

BASELINE_CSV = ROOT / "results" / "chapter5_q11_results" / "road_baseline_solution_comparison.csv"
ALGO_CSV = ROOT / "results" / "chapter5_algorithm_comparison" / "algorithm_performance_summary.csv"
ABLATION_CSV = ROOT / "results" / "chapter5_ablation_experiment_5.16" / "results" / "ablation_summary.csv"
WEIGHT_SCENARIO_CSV = ROOT / "results" / "5.20" / "weight_scenario_topsis_final_solutions.csv"


def set_cn_font() -> None:
    plt.rcParams["font.sans-serif"] = [
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 120


def fmt(value: float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}"


def fmt_pct(value: float, digits: int = 2, signed: bool = True) -> str:
    sign = "+" if signed and float(value) > 0 else ""
    return f"{sign}{float(value):.{digits}f}%"


def cell_text(cell, text: str, font_size: int = 9) -> None:
    cell.text = str(text)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.name = "宋体"


def style_table(table, font_size: int = 9) -> None:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = "宋体"


def set_table_data(table, headers: list[str], rows: list[list[str]]) -> None:
    target_rows = len(rows) + 1
    while len(table.rows) < target_rows:
        table.add_row()
    while len(table.rows) > target_rows:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    for j, h in enumerate(headers):
        cell_text(table.cell(0, j), h)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell_text(table.cell(i, j), value)
    style_table(table)


def paragraph_after(paragraph, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = paragraph._parent.add_paragraph()
    p._p = new_p
    if style is not None:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_table_after(paragraph, doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_data(table, headers, rows)
    paragraph._p.addnext(table._tbl)
    return table


def insert_picture_after(paragraph, image_path: Path, caption: str, width_cm: float = 15.0):
    cap = paragraph_after(paragraph, caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.font.name = "宋体"
        run.font.size = Pt(10)
    pic_p = paragraph_after(cap)
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    return pic_p


def find_paragraph(doc: Document, keyword: str):
    for p in doc.paragraphs:
        if keyword in p.text:
            return p
    return None


def find_table(doc: Document, required_keywords: list[str]):
    for table in doc.tables:
        text = "\n".join(cell.text for row in table.rows for cell in row.cells)
        if all(k in text for k in required_keywords):
            return table
    return None


def make_baseline_figure(df: pd.DataFrame) -> Path:
    labels = ["全公路", "最小成本", "TOPSIS折中", "最大服务"]
    cost = df["total_cost_yuan_day"].to_numpy()
    service = df["metro_service_ratio"].to_numpy()
    direct = df["direct_volume"].to_numpy()
    carbon = df["total_carbon_tCO2"].to_numpy()

    fig, axes = plt.subplots(1, 2, figsize=(11.8, 4.6))
    x = np.arange(len(labels))
    axes[0].bar(x, cost, color="#6b8fb3", edgecolor="black", linewidth=0.6, label="综合成本")
    axes[0].set_ylabel("综合成本（元/日）")
    axes[0].set_xticks(x)
    axes[0].set_xticklabels(labels)
    axes[0].grid(axis="y", linestyle="--", alpha=0.35)
    ax2 = axes[0].twinx()
    ax2.plot(x, service, marker="o", color="#b45f5f", linewidth=2.0, label="地铁服务比例")
    ax2.set_ylabel("地铁服务比例")
    lines1, labels1 = axes[0].get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    axes[0].legend(lines1 + lines2, labels1 + labels2, loc="upper left", frameon=False)

    width = 0.36
    axes[1].bar(x - width / 2, direct, width, color="#84a98c", edgecolor="black", linewidth=0.6, label="公路直送量")
    axes[1].set_ylabel("公路直送量（SLU/日）")
    axes[1].set_xticks(x)
    axes[1].set_xticklabels(labels)
    axes[1].grid(axis="y", linestyle="--", alpha=0.35)
    ax3 = axes[1].twinx()
    ax3.bar(x + width / 2, carbon, width, color="#c9a66b", edgecolor="black", linewidth=0.6, label="碳排放")
    ax3.set_ylabel("碳排放（tCO2/日）")
    h1, l1 = axes[1].get_legend_handles_labels()
    h2, l2 = ax3.get_legend_handles_labels()
    axes[1].legend(h1 + h2, l1 + l2, loc="upper right", frameon=False)

    for ax in axes:
        for spine in ax.spines.values():
            spine.set_visible(True)
    fig.tight_layout()
    out = OUT_DIR / "fig_5_scheme_baseline_comparison.png"
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return out


def make_algorithm_figure(df: pd.DataFrame) -> Path:
    metrics = [
        ("HV_mean", "HV（越大越优）"),
        ("IGD_mean", "IGD（越小越优）"),
        ("Spacing_mean", "Spacing（越小越优）"),
        ("Spread_mean", "Spread"),
    ]
    labels = df["algorithm_label"].tolist()
    colors = ["#587792", "#8f6f56", "#7c9a72", "#a45d63"]
    fig, axes = plt.subplots(2, 2, figsize=(10.5, 7.0))
    for ax, (col, ylabel) in zip(axes.flat, metrics):
        values = df[col].to_numpy()
        ax.bar(labels, values, color=colors, edgecolor="black", linewidth=0.6)
        ax.set_ylabel(ylabel)
        ax.tick_params(axis="x", rotation=15)
        ax.grid(axis="y", linestyle="--", alpha=0.35)
        for spine in ax.spines.values():
            spine.set_visible(True)
    fig.tight_layout()
    out = OUT_DIR / "fig_5_algorithm_performance_comparison.png"
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return out


def make_ablation_figure(df: pd.DataFrame) -> Path:
    metrics = [
        ("HV_mean", "HV（越大越优）"),
        ("IGD_mean", "IGD（越小越优）"),
        ("Spacing_mean", "Spacing（越小越优）"),
        ("Spread_mean", "Spread"),
    ]
    labels = df["algorithm"].tolist()
    colors = ["#587792", "#8f6f56", "#7c9a72", "#a45d63"]
    fig, axes = plt.subplots(2, 2, figsize=(10.5, 7.2))
    for ax, (col, ylabel) in zip(axes.flat, metrics):
        values = df[col].to_numpy()
        ax.bar(labels, values, color=colors, edgecolor="black", linewidth=0.6)
        ax.set_ylabel(ylabel)
        ax.tick_params(axis="x", rotation=18)
        ax.grid(axis="y", linestyle="--", alpha=0.35)
        for spine in ax.spines.values():
            spine.set_visible(True)
    fig.tight_layout()
    out = OUT_DIR / "fig_5_ablation_metrics_comparison.png"
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return out


def main() -> None:
    set_cn_font()
    baseline = pd.read_csv(BASELINE_CSV)
    algo = pd.read_csv(ALGO_CSV)
    ablation = pd.read_csv(ABLATION_CSV)

    baseline_fig = make_baseline_figure(baseline)
    algo_fig = make_algorithm_figure(algo)
    ablation_fig = make_ablation_figure(ablation)

    doc = Document(TARGET)

    # Replace early uncertainty note now that the requested data have been checked.
    p_note = find_paragraph(doc, "本章不虚构实验结果")
    if p_note is not None:
        p_note.text = (
            "本章数据均依据项目数据目录中的模型输入、算法输出和实验结果文件进行整理。"
            "对于当前结果文件能够核验的指标，已在表格中给出具体数值；对于仍需与统计年鉴页码、企业原始台账或图件编号进一步核对的内容，保留“待核验”说明。"
        )

    p_weight = find_paragraph(doc, "为增强方案解释的完整性，第五章可进一步设置")
    if p_weight is not None:
        p_weight.text = (
            "在方案后评价中，若决策主体未给出明确偏好，本文采用归一化等权作为基准设置，"
            "即成本目标和地铁服务比例目标具有相同权重。为增强方案解释的完整性，本章进一步设置成本偏好、均衡偏好和服务偏好三类权重情景，"
            "其作用仅限于对既有 Pareto 非支配方案进行 TOPSIS 后评价，不改变第四章双目标布局优化模型本身。"
        )

    # Table 1: baseline comparison.
    baseline_names = {
        "all_road": "全公路直送",
        "min_cost": "最小成本方案",
        "topsis_compromise": "TOPSIS 折中方案",
        "max_service": "最大服务方案",
    }
    baseline_rows = []
    for _, r in baseline.iterrows():
        scenario = baseline_names.get(r["scenario_en"], r["scenario_en"])
        baseline_rows.append(
            [
                scenario,
                fmt(r["total_cost_yuan_day"], 2),
                fmt(r["metro_service_ratio"], 6) if r["metro_service_ratio"] else "0",
                fmt(r["direct_volume"], 2),
                fmt(r["total_carbon_tCO2"], 6),
                "0" if abs(r["total_cost_change_pct"]) < 1e-12 else fmt_pct(r["total_cost_change_pct"], 6),
            ]
        )
    t_baseline = find_table(doc, ["方案", "综合成本", "地铁服务比例", "相对全公路综合成本变化"])
    if t_baseline is None:
        raise RuntimeError("Cannot find baseline comparison table.")
    set_table_data(
        t_baseline,
        ["方案", "综合成本", "地铁服务比例", "公路直送量", "碳排放（tCO2/day）", "相对全公路综合成本变化"],
        baseline_rows,
    )

    # Insert/update baseline figure after baseline interpretation paragraph.
    p_baseline = find_paragraph(doc, "全公路方案综合成本为 204130.69")
    if p_baseline is not None and "图5-" not in p_baseline.text:
        insert_picture_after(
            p_baseline,
            baseline_fig,
            "图5-X 不同布局方案与全公路基准方案的成本、服务比例、直送量和碳排放对比",
        )

    # Table 2: algorithm comparison.
    algo_rows = []
    for _, r in algo.iterrows():
        algo_rows.append(
            [
                r["algorithm_label"],
                fmt(r["HV_mean"], 6),
                fmt(r["IGD_mean"], 6),
                fmt(r["Spacing_mean"], 6),
                fmt(r["Spread_mean"], 6),
                fmt(r["unique_solutions_mean"], 1),
                fmt(r["runtime_seconds_mean"], 2),
            ]
        )
    t_algo = find_table(doc, ["算法", "HV", "IGD", "Spacing", "Spread", "唯一解均值"])
    if t_algo is None:
        raise RuntimeError("Cannot find algorithm comparison table.")
    set_table_data(t_algo, ["算法", "HV", "IGD", "Spacing", "Spread", "唯一解均值", "运行时间（s）"], algo_rows)

    p_algo = find_paragraph(doc, "QVNS-NSGA-II 在解集质量方面优于典型多目标优化算法")
    if p_algo is not None:
        p_algo.text = (
            "算法对比结果表明，QVNS-NSGA-II 在 HV 和 IGD 两项解集质量指标上优于典型多目标优化算法，"
            "但相较标准 NSGA-II 存在一定计算时间增加。MOPSO 与 SPEA2 的运行时间相对较长，且解集收敛性和综合分布质量低于 QVNS-NSGA-II。"
            "因此，本文算法的优势主要体现在解集质量和搜索稳定性的综合平衡，而非单纯追求最短运行时间。"
        )
        insert_picture_after(p_algo, algo_fig, "图5-X 典型多目标优化算法性能指标对比")

    # Insert ablation table and figure near the ablation paragraph.
    p_ablation = find_paragraph(doc, "为进一步说明算法机制作用")
    if p_ablation is not None:
        p_ablation.text = (
            "为进一步说明算法机制作用，本章设置消融实验，对比 QVNS-Q11C、VNS-only、Q-learning-only 和标准 NSGA-II 等算法变体。"
            "消融实验重点检验变邻域搜索、Q-learning 动作选择与档案保护机制对解集质量和搜索稳定性的影响。"
            "需要谨慎指出，QVNS-Q11C 与 VNS-only 在部分指标上较为接近，不能简单断言组合机制在所有指标上均绝对最优；"
            "更稳妥的结论是，相较于标准 NSGA-II 和单独 Q-learning 机制，融合变邻域搜索、Q-learning 动作选择与档案保护机制的算法在解集质量、搜索稳定性和机制可解释性方面表现更均衡。"
        )

        existing_ablation_table = find_table(doc, ["QVNS-Q11C", "VNS-only", "Q-learning-only"])
        if existing_ablation_table is None:
            ablation_rows = []
            for _, r in ablation.iterrows():
                ablation_rows.append(
                    [
                        r["algorithm"],
                        fmt(r["HV_mean"], 6),
                        fmt(r["IGD_mean"], 6),
                        fmt(r["Spacing_mean"], 6),
                        fmt(r["Spread_mean"], 6),
                        fmt(r["unique_points_mean"], 1),
                        fmt(r["runtime_seconds_mean"], 2),
                        fmt(r["safeguard_rejects_mean"], 1),
                        fmt(r["actual_vns_calls_mean"], 1),
                    ]
                )
            insert_table_after(
                p_ablation,
                doc,
                ["算法变体", "HV", "IGD", "Spacing", "Spread", "唯一解均值", "运行时间（s）", "档案保护拒绝次数", "VNS 调用次数"],
                ablation_rows,
            )
        insert_picture_after(p_ablation, ablation_fig, "图5-X QVNS-NSGA-II 消融实验性能指标对比")

    # Add weight scenario final scheme table after the weight table paragraph, if available.
    if WEIGHT_SCENARIO_CSV.exists() and find_table(doc, ["权重情景", "解编号", "TOPSIS贴近度"]) is None:
        ws = pd.read_csv(WEIGHT_SCENARIO_CSV, encoding="utf-8-sig")
        rows = []
        for _, r in ws.iterrows():
            rows.append(
                [
                    r["权重情景"],
                    fmt(r["Z1综合成本权重"], 2),
                    fmt(r["Z2地铁服务比例权重"], 2),
                    str(int(r["solution_id"])),
                    fmt(r["TOPSIS贴近度"], 6),
                    fmt(r["Z1_total_cost"], 2),
                    fmt(r["Z2_metro_service_ratio"], 6),
                    str(int(r["num_open"])),
                    fmt(r["metro_volume"], 2),
                    fmt(r["direct_volume"], 2),
                ]
            )
        p_ws = find_paragraph(doc, "不改变第四章双目标布局优化模型本身")
        if p_ws is not None:
            insert_table_after(
                p_ws,
                doc,
                ["权重情景", "$Z_1$ 权重", "$Z_2$ 权重", "解编号", "TOPSIS贴近度", "综合成本", "地铁服务比例", "启用数", "地铁联运量", "公路直送量"],
                rows,
            )

    try:
        doc.save(TARGET)
        saved_to = TARGET
    except PermissionError:
        doc.save(FALLBACK_TARGET)
        saved_to = FALLBACK_TARGET
        print(f"target_locked={TARGET}")
    print(f"updated_doc={saved_to}")
    print(f"baseline_fig={baseline_fig}")
    print(f"algorithm_fig={algo_fig}")
    print(f"ablation_fig={ablation_fig}")


if __name__ == "__main__":
    main()
