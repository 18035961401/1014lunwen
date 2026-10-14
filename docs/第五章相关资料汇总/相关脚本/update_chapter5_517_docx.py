from __future__ import annotations

import shutil
import tempfile
import os
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "docs" / "同步更新输出_5.17"
FIG_DIR = DOC_DIR / "nature_figures_q11_runs10_v2"
RESULT_DIR = ROOT / "results" / "chapter5_q11_results"
DOCX_PATH = DOC_DIR / "第五章_同步更新版.docx"
BACKUP_PATH = DOC_DIR / "第五章_同步更新版_before_fig_sensitivity_rewrite.docx"


def insert_paragraph_after(anchor, text: str = "", bold: bool = False, align=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(10.5)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def insert_paragraph_after_element(doc: Document, anchor_el, text: str = "", bold: bool = False, align=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    paragraph = Paragraph(new_p, doc._body)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(10.5)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def move_block_after(anchor_el, block_el):
    parent = block_el.getparent()
    if parent is not None:
        parent.remove(block_el)
    anchor_el.addnext(block_el)
    return block_el


def insert_table_after(doc: Document, anchor: Paragraph, data: list[list[str]]) -> Table:
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.style = "Table Grid"
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    if r == 0:
                        run.bold = True
    move_block_after(anchor._p, table._tbl)
    return table


def insert_picture_after(doc: Document, anchor: Paragraph, image_path: Path, width_in: float) -> Paragraph:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))
    move_block_after(anchor._p, paragraph._p)
    return paragraph


def remove_between(doc: Document, start_text: str, end_text: str) -> tuple[Paragraph, Paragraph]:
    start = next(p for p in doc.paragraphs if p.text.strip() == start_text)
    end = next(p for p in doc.paragraphs if p.text.strip() == end_text)
    body = doc.element.body
    children = list(body)
    start_i = children.index(start._p)
    end_i = children.index(end._p)
    for el in children[start_i + 1 : end_i]:
        body.remove(el)
    return start, end


def fmt_money(x: float) -> str:
    return f"{x:,.2f}"


def fmt_pct(x: float) -> str:
    return f"{x * 100:.2f}%"


def build_sensitivity_table(summary: pd.DataFrame) -> list[list[str]]:
    type_map = {"demand_scale": "需求规模", "service_radius": "服务半径", "source_share": "货源比例"}
    order = [
        ("demand_scale", "low"),
        ("demand_scale", "base"),
        ("demand_scale", "high"),
        ("service_radius", "r3"),
        ("service_radius", "r4"),
        ("service_radius", "r5"),
        ("source_share", "s5_5"),
        ("source_share", "s6_4"),
        ("source_share", "s7_3"),
        ("source_share", "s8_2"),
        ("source_share", "s9_1"),
    ]
    rows = [["情景类型", "情景", "运行次数", "代表HV", "综合成本/元·日^-1", "地铁服务比例", "启用站点"]]
    for group, case in order:
        row = summary[(summary["group_type"] == group) & (summary["case_key"] == case)].iloc[0]
        rows.append(
            [
                type_map[group],
                case,
                str(int(row["runs"])),
                f"{row['representative_HV']:.4f}",
                fmt_money(row["topsis_cost"]),
                fmt_pct(row["topsis_service_ratio"]),
                str(row["topsis_open_stations"]),
            ]
        )
    return rows


def paragraph_texts(summary: pd.DataFrame) -> dict[str, str]:
    def row(group: str, case: str):
        return summary[(summary["group_type"] == group) & (summary["case_key"] == case)].iloc[0]

    low, base, high = row("demand_scale", "low"), row("demand_scale", "base"), row("demand_scale", "high")
    r3, r4, r5 = row("service_radius", "r3"), row("service_radius", "r4"), row("service_radius", "r5")
    s55, s64, s73, s82, s91 = [row("source_share", key) for key in ["s5_5", "s6_4", "s7_3", "s8_2", "s9_1"]]
    return {
        "intro": "为避免单一综合图弱化空间机理解释，本文在统一 p100-g300-runs10 口径下，将参数情景敏感性分析进一步拆分为需求规模、服务半径和物流货源比例三类情景。每个独立情景均运行 10 次，并选取 HV 最大的运行作为代表运行；基准情景仅求解一次，并同步用于 demand_scale/base、service_radius/r4、source_share/s8_2 与主方案目录，以保证同一基准口径下结果一致。",
        "demand": f"需求规模变化主要体现为系统总量约束对成本和布局稳定性的影响。低需求、基准和高需求情景的折中解成本分别为 {fmt_money(low.topsis_cost)}、{fmt_money(base.topsis_cost)} 和 {fmt_money(high.topsis_cost)} 元/日，随需求强度提高呈递增趋势。空间布局上，低需求情景启用 {int(low.topsis_open_count)} 个站点（{low.topsis_open_stations}），而基准与高需求情景均启用 {int(base.topsis_open_count)} 个站点（{base.topsis_open_stations}），说明火车南站—高升桥—省体育馆—太平园组合在中高需求下具有较强的结构稳定性。值得注意的是，高需求情景的地铁服务比例为 {fmt_pct(high.topsis_service_ratio)}，低于基准情景的 {fmt_pct(base.topsis_service_ratio)}，表明当需求总量增加后，部分边缘需求仍需由公路直送承担，地铁联运能力提升并非简单随需求规模线性增加。",
        "radius": f"服务半径变化直接改变末端需求点可被候选转运点吸收的空间范围。3 km、4 km 和 5 km 情景下，折中解成本分别为 {fmt_money(r3.topsis_cost)}、{fmt_money(r4.topsis_cost)} 和 {fmt_money(r5.topsis_cost)} 元/日；其中 4 km 基准情景与主方案完全一致，服务比例达到 {fmt_pct(r4.topsis_service_ratio)}。当半径收缩至 3 km 时，模型保留桐梓林以弥补南部覆盖不足，但服务比例降至 {fmt_pct(r3.topsis_service_ratio)}；当半径放宽至 5 km 时，模型减少为 {int(r5.topsis_open_count)} 个启用站点（{r5.topsis_open_stations}），成本降至 {fmt_money(r5.topsis_cost)} 元/日，但服务比例仅为 {fmt_pct(r5.topsis_service_ratio)}。这说明服务半径扩张会提高空间可达性，却也可能诱导模型选择更少站点，从而降低联运覆盖强度。",
        "source": f"物流货源比例变化对布局结果的扰动最为明显，反映了上游货源空间结构对联运组织强度的外生影响。当货源比例由 5:5 调整至 9:1 时，折中解成本由 {fmt_money(s55.topsis_cost)} 元/日降至 {fmt_money(s91.topsis_cost)} 元/日，地铁服务比例由 {fmt_pct(s55.topsis_service_ratio)} 降至 {fmt_pct(s91.topsis_service_ratio)}。5:5 与 6:4 情景分别启用 {int(s55.topsis_open_count)} 和 {int(s64.topsis_open_count)} 个站点，空间上向神仙树、孵化园、金融城等南部与高新区方向扩展；而 9:1 情景仅启用 {int(s91.topsis_open_count)} 个站点（{s91.topsis_open_stations}），更强调低成本路径。由此可见，货源比例不仅改变成本规模，更会改变启用站点的空间重心，是影响地铁物流联运强度和末端接驳组织方式的关键参数。",
    }


def update_docx_body() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)
    doc = Document(DOCX_PATH)
    summary = pd.read_csv(RESULT_DIR / "q11_sensitivity_summary.csv")
    texts = paragraph_texts(summary)
    start, _ = remove_between(doc, "5.4.3 参数情景敏感性分析", "5.5 本章小结")
    anchor = start
    anchor = insert_paragraph_after(anchor, texts["intro"])
    anchor = insert_paragraph_after(anchor, "表 5-14 Q11C 敏感性分析代表性折中解汇总", align=WD_ALIGN_PARAGRAPH.CENTER)
    table = insert_table_after(doc, anchor, build_sensitivity_table(summary))

    anchor = insert_paragraph_after_element(doc, table._tbl, "5.4.3.1 需求规模变化下的布局结果分析", bold=True)
    anchor = insert_paragraph_after(anchor, texts["demand"])
    anchor = insert_picture_after(doc, anchor, FIG_DIR / "fig5_13_demand_scale_sensitivity_map.png", 6.25)
    anchor = insert_paragraph_after(anchor, "图 5-13 需求规模变化下的布局结果与成本—服务响应", align=WD_ALIGN_PARAGRAPH.CENTER)

    anchor = insert_paragraph_after(anchor, "5.4.3.2 服务半径变化下的布局结果分析", bold=True)
    anchor = insert_paragraph_after(anchor, texts["radius"])
    anchor = insert_picture_after(doc, anchor, FIG_DIR / "fig5_14_service_radius_sensitivity_map.png", 6.25)
    anchor = insert_paragraph_after(anchor, "图 5-14 服务半径变化下的布局结果与成本—服务响应", align=WD_ALIGN_PARAGRAPH.CENTER)

    anchor = insert_paragraph_after(anchor, "5.4.3.3 物流货源比例变化下的布局结果分析", bold=True)
    anchor = insert_paragraph_after(anchor, texts["source"])
    anchor = insert_picture_after(doc, anchor, FIG_DIR / "fig5_15_source_share_sensitivity_map.png", 6.25)
    anchor = insert_paragraph_after(anchor, "图 5-15 物流货源比例变化下的布局结果与成本—服务响应", align=WD_ALIGN_PARAGRAPH.CENTER)

    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("优化结果表明，武侯区地铁物流布局存在明确的成本—服务权衡关系。"):
            paragraph.text = (
                "优化结果表明，武侯区地铁物流布局存在明确的成本—服务权衡关系。按 p100-g300-runs10 口径选取 HV 最大的基准代表运行后，TOPSIS 折中方案启用 4 个转运点（火车南站、高升桥、省体育馆、太平园），综合成本为 207,005.43 元/日，地铁服务比例为 42.39%。与全公路基准相比，折中方案虽然综合成本略有增加，但公路相关成本和碳排放明显下降，说明地铁联运的价值主要体现在道路负荷缓解和低碳绩效改善。消融实验进一步表明，QVNS-Q11C 在 HV、IGD、Spacing 与机制触发稳定性方面具有综合优势，Q-learning、VNS 与 safeguard 机制分别承担动作自适应、局部强化搜索和不可行退化抑制作用。敏感性分析显示，需求规模提高会推高系统成本，服务半径改变会影响站点组合与覆盖结构，物流货源比例则会显著改变启用站点空间重心和联运服务强度。整体来看，地铁物流转运点布局应优先关注火车南站、太平园、省体育馆、高升桥等兼具网络连接性、需求邻近性和负载可控性的枢纽或换乘站点，并结合平峰期班次能力、末端服务半径和货源组织结构进行动态调整。"
            )
            break
    doc.save(DOCX_PATH)


def replace_media(docx_path: Path) -> None:
    replacements = {
        "word/media/image4.png": FIG_DIR / "fig5_4_critic_weight_pie.png",
        "word/media/image6.png": FIG_DIR / "fig5_6_qvns_pareto_front_updated.png",
        "word/media/image10.png": FIG_DIR / "fig5_10_qlearning_safeguard_updated.png",
        "word/media/image12.png": FIG_DIR / "fig5_12_ablation_mechanism_updated.png",
    }
    fd, tmp_name = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    tmp = Path(tmp_name)
    with ZipFile(docx_path, "r") as zin, ZipFile(tmp, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in replacements:
                data = replacements[item.filename].read_bytes()
            zout.writestr(item, data)
    shutil.move(str(tmp), docx_path)


def main() -> None:
    update_docx_body()
    replace_media(DOCX_PATH)
    print(f"Updated: {DOCX_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
