# -*- coding: utf-8 -*-
"""Build the complete Chapter 5 Word document from optimized Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = PROJECT_ROOT / "results" / "chapter5_draft.md"
OUTPUT_DOCX = PROJECT_ROOT / "data" / "第五章_成都市武侯区地铁物流配送转运点布局优化研究.docx"


def set_run_font(run, east_asia="宋体", ascii_font="Times New Roman", size=10.5, bold=False):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold=False, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text.strip())
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(3)

    for style_name, east_font, ascii_font, size, bold in [
        ("Heading 1", "黑体", "Times New Roman", 16, True),
        ("Heading 2", "黑体", "Times New Roman", 14, True),
        ("Heading 3", "黑体", "Times New Roman", 12, True),
    ]:
        style = styles[style_name]
        style.font.name = ascii_font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25


def strip_bold(text: str) -> str:
    return text.replace("**", "").strip()


def add_text_paragraph(doc: Document, text: str):
    text = text.strip()
    if not text:
        return
    paragraph = doc.add_paragraph()
    if text.startswith(">"):
        text = text.lstrip("> ").strip()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.right_indent = Cm(0.5)
    if text.startswith("**") and text.endswith("**") and (text.startswith("**表") or text.startswith("**图")):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(strip_bold(text))
        set_run_font(run, east_asia="黑体", size=10.5, bold=True)
        return
    if text.startswith("\\[") or text.startswith("\\]") or text.startswith("\\"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(text)
        set_run_font(run, east_asia="宋体", ascii_font="Cambria Math", size=10)
        return
    run = paragraph.add_run(strip_bold(text))
    set_run_font(run)


def split_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def add_markdown_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    n_cols = max(len(row) for row in rows)
    rows = [row + [""] * (n_cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    font_size = 7.5 if n_cols >= 8 else 8.5 if n_cols >= 6 else 9
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_text(cell, strip_bold(text), bold=(r_idx == 0), size=font_size)
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
    doc.add_paragraph()


def resolve_image_path(path_text: str) -> Path:
    path = Path(path_text.strip())
    if path.is_absolute():
        return path
    candidate = PROJECT_ROOT / path
    if candidate.exists():
        return candidate
    candidate = SOURCE_MD.parent / path
    if candidate.exists():
        return candidate
    return PROJECT_ROOT / path


def add_image(doc: Document, alt: str, path_text: str):
    image_path = resolve_image_path(path_text)
    if not image_path.exists():
        add_text_paragraph(doc, f"[图片缺失：{alt} - {image_path}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    width = Cm(15.5)
    if "fig_compromise_map" in image_path.name:
        width = Cm(16.2)
    elif "quality_metrics" in image_path.name or "boxplots" in image_path.name:
        width = Cm(16.0)
    run.add_picture(str(image_path), width=width)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption_run = caption.add_run(alt.strip())
    set_run_font(caption_run, east_asia="黑体", size=9.5, bold=True)


def build_docx():
    text = SOURCE_MD.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    doc = Document()
    configure_document(doc)
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_heading(line[2:].strip(), level=1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if image_match:
            add_image(doc, image_match.group(1), image_match.group(2))
            i += 1
            continue
        if line.strip().startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_markdown_table(doc, rows)
            continue
        add_text_paragraph(doc, line)
        i += 1
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    print(build_docx())
